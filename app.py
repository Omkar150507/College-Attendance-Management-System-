from flask import Flask, render_template, request, redirect, url_for, session, flash, Response
import os
import sqlite3
from urllib.parse import urlparse
from xml.sax.saxutils import escape as xml_escape

try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
    from psycopg2.pool import ThreadedConnectionPool
    POSTGRES_AVAILABLE = True
except ImportError:
    psycopg2 = None
    RealDictCursor = None
    ThreadedConnectionPool = None
    POSTGRES_AVAILABLE = False

from flask import Flask, render_template, request, redirect, url_for, session, flash, Response
from pathlib import Path
from functools import wraps
from datetime import date, datetime
from zoneinfo import ZoneInfo
import csv
import io
import json
import zipfile
import smtplib
import re
import secrets
from calendar import monthrange
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.drawing.image import Image as XLImage
    OPENPYXL_AVAILABLE = True
except ImportError:
    Workbook = None
    XLImage = None
    OPENPYXL_AVAILABLE = False
from email.message import EmailMessage
try:
    import qrcode
    QRCODE_AVAILABLE = True
except ImportError:
    qrcode = None
    QRCODE_AVAILABLE = False
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix
from config import COLLEGE_NAME, COLLEGE_TAGLINE, LOGO_FILE

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "attendance.db"
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
RESET_TOKEN_MAX_AGE = 30 * 60

# Small per-worker PostgreSQL pool. Render can occasionally recycle an idle
# PostgreSQL/TLS connection. A pooled connection that has gone stale must never
# be handed back to a request; otherwise pages can intermittently return HTTP 500
# with errors such as: "SSL error: decryption failed or bad record mac".
PG_POOL = None

def _discard_pg_connection(conn):
    """Return a bad PostgreSQL connection to the pool and force it to close."""
    if conn is None or PG_POOL is None:
        return
    try:
        PG_POOL.putconn(conn, close=True)
    except Exception:
        try:
            conn.close()
        except Exception:
            pass

def _new_pg_pool():
    if not POSTGRES_AVAILABLE:
        raise RuntimeError("PostgreSQL driver is not installed. Run: pip install psycopg2-binary")
    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL is not configured")

    # TCP keepalives reduce the chance that an idle Render/PostgreSQL connection
    # survives in the pool after its network/TLS session has already expired.
    return ThreadedConnectionPool(
        1,
        int(os.environ.get("DB_POOL_MAX", "10")),
        DATABASE_URL,
        sslmode=os.environ.get("DB_SSLMODE", "require"),
        connect_timeout=10,
        keepalives=1,
        keepalives_idle=int(os.environ.get("DB_KEEPALIVES_IDLE", "30")),
        keepalives_interval=int(os.environ.get("DB_KEEPALIVES_INTERVAL", "10")),
        keepalives_count=int(os.environ.get("DB_KEEPALIVES_COUNT", "5")),
    )

def get_pg_pool():
    global PG_POOL
    if PG_POOL is None:
        PG_POOL = _new_pg_pool()
    return PG_POOL

def get_healthy_pg_connection():
    """Get a live connection; discard stale/broken pooled connections."""
    pool = get_pg_pool()
    last_error = None
    for _ in range(3):
        conn = pool.getconn()
        try:
            if getattr(conn, "closed", 0):
                _discard_pg_connection(conn)
                continue
            # A lightweight round-trip detects stale TLS/socket state before the
            # actual application query uses the connection.
            with conn.cursor() as cur:
                cur.execute("SELECT 1")
                cur.fetchone()
            conn.rollback()
            conn.autocommit = False
            return conn
        except Exception as exc:
            last_error = exc
            _discard_pg_connection(conn)
    raise RuntimeError(f"Unable to obtain a healthy PostgreSQL connection: {last_error}")

def get_smtp_settings():
    """Read SMTP settings at request time so Render environment changes are picked up after restart."""
    host = os.environ.get("SMTP_HOST", "").strip()
    username = os.environ.get("SMTP_USERNAME", "").strip()
    password = os.environ.get("SMTP_PASSWORD", "").strip().replace(" ", "")
    from_email = os.environ.get("SMTP_FROM_EMAIL", "").strip() or username
    try:
        port = int(os.environ.get("SMTP_PORT", "587") or 587)
    except ValueError:
        port = 587
    use_tls = os.environ.get("SMTP_USE_TLS", "1").strip().lower() not in ("0", "false", "no")
    return host, port, username, password, from_email, use_tls
YEAR_OPTIONS = ["1st Year", "2nd Year", "3rd Year", "4th Year"]
SECURITY_QUESTIONS = [
    "What was the name of your first school?",
    "What is the name of your hometown?",
    "What was your favorite subject in school?",
    "What is your favorite teacher's name?",
]

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
app.secret_key = os.environ.get("SECRET_KEY", "dev-only-change-this-secret-key")
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=bool(DATABASE_URL),
)

def reset_serializer():
    return URLSafeTimedSerializer(app.secret_key, salt="password-reset")


def send_password_reset_email(to_email, full_name, reset_link):
    host, port, username, password, from_email, use_tls = get_smtp_settings()
    missing = []
    if not host: missing.append("SMTP_HOST")
    if not username: missing.append("SMTP_USERNAME")
    if not password: missing.append("SMTP_PASSWORD")
    if not from_email: missing.append("SMTP_FROM_EMAIL")
    if missing:
        raise RuntimeError("Password reset email is not configured. Missing: " + ", ".join(missing))
    app.logger.info("SMTP configuration detected: host=%s port=%s username=%s from_email=%s tls=%s", host, port, username, from_email, use_tls)
    msg = EmailMessage()
    msg["Subject"] = f"{COLLEGE_NAME} - Password Reset"
    msg["From"] = from_email
    msg["To"] = to_email
    msg.set_content(
        f"Hello {full_name or 'User'},\n\n"
        f"A password reset was requested for your {COLLEGE_NAME} account.\n\n"
        f"Open this link within 30 minutes to set a new password:\n{reset_link}\n\n"
        "If you did not request this, you can ignore this email.\n"
    )
    if use_tls and port == 465:
        with smtplib.SMTP_SSL(host, port, timeout=20) as server:
            server.login(username, password)
            server.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=20) as server:
            if use_tls:
                server.starttls()
            server.login(username, password)
            server.send_message(msg)


class CompatRow(dict):
    def __getitem__(self, key):
        if isinstance(key, int):
            return list(self.values())[key]
        return super().__getitem__(key)

class CompatCursor:
    def __init__(self, cursor, postgres=False):
        self.cursor = cursor
        self.postgres = postgres
        self.lastrowid = None

    def fetchone(self):
        row = self.cursor.fetchone()
        return CompatRow(row) if row is not None else None

    def fetchall(self):
        return [CompatRow(r) for r in self.cursor.fetchall()]

class CompatDB:
    def __init__(self):
        self.postgres = bool(DATABASE_URL)
        self.conn = None
        self.pool = None
        self._closed = False
        if self.postgres:
            self.pool = get_pg_pool()
            self.conn = get_healthy_pg_connection()
            self.conn.autocommit = False
        else:
            self.conn = sqlite3.connect(DB_PATH)
            self.conn.row_factory = sqlite3.Row
            self.conn.execute("PRAGMA foreign_keys = ON")

    def _prepare_sql(self, sql):
        sql = sql.replace("?", "%s")
        # SQLite-specific INSERT OR IGNORE -> PostgreSQL equivalent.
        sql = sql.replace("INSERT OR IGNORE INTO", "INSERT INTO")
        if "INSERT INTO subject_teachers" in sql and "ON CONFLICT" not in sql:
            sql = sql.rstrip().rstrip(";") + " ON CONFLICT (teacher_id, subject_id) DO NOTHING"
        return sql

    def _replace_broken_connection(self):
        old = self.conn
        _discard_pg_connection(old)
        self.conn = get_healthy_pg_connection()
        self.conn.autocommit = False

    def execute(self, sql, params=()):
        if self.postgres:
            sql = self._prepare_sql(sql)
            try:
                cur = self.conn.cursor(cursor_factory=RealDictCursor)
                cur.execute(sql, params)
                return CompatCursor(cur, True)
            except (psycopg2.OperationalError, psycopg2.InterfaceError) as exc:
                # Retry only when the connection is still at the beginning of a
                # transaction. Retrying inside a partially completed transaction
                # could silently lose earlier writes. This safely fixes the common
                # stale pooled-connection case seen on read-only page loads.
                try:
                    can_retry = self.conn.get_transaction_status() == psycopg2.extensions.TRANSACTION_STATUS_IDLE
                except Exception:
                    can_retry = False
                if not can_retry:
                    raise exc
                self._replace_broken_connection()
                cur = self.conn.cursor(cursor_factory=RealDictCursor)
                cur.execute(sql, params)
                return CompatCursor(cur, True)
        cur = self.conn.execute(sql, params)
        wrapped = CompatCursor(cur, False)
        wrapped.lastrowid = cur.lastrowid
        return wrapped

    def executemany(self, sql, seq_of_params):
        """Execute a parameterized statement for each parameter tuple."""
        if self.postgres:
            sql = self._prepare_sql(sql)
            try:
                cur = self.conn.cursor(cursor_factory=RealDictCursor)
                cur.executemany(sql, seq_of_params)
                return CompatCursor(cur, True)
            except (psycopg2.OperationalError, psycopg2.InterfaceError) as exc:
                try:
                    can_retry = self.conn.get_transaction_status() == psycopg2.extensions.TRANSACTION_STATUS_IDLE
                except Exception:
                    can_retry = False
                if not can_retry:
                    raise exc
                self._replace_broken_connection()
                cur = self.conn.cursor(cursor_factory=RealDictCursor)
                cur.executemany(sql, seq_of_params)
                return CompatCursor(cur, True)
        cur = self.conn.executemany(sql, seq_of_params)
        wrapped = CompatCursor(cur, False)
        wrapped.lastrowid = cur.lastrowid
        return wrapped

    def executescript(self, script):
        if self.postgres:
            # PostgreSQL schema used for hosted deployment.
            statements = [x.strip() for x in script.split(";") if x.strip()]
            for stmt in statements:
                stmt = stmt.replace("INTEGER PRIMARY KEY AUTOINCREMENT", "SERIAL PRIMARY KEY")
                stmt = stmt.replace("CURRENT_TIMESTAMP", "CURRENT_TIMESTAMP")
                stmt = stmt.replace("TEXT DEFAULT ''", "TEXT DEFAULT ''")
                self.execute(stmt)
        else:
            self.conn.executescript(script)

    def commit(self):
        self.conn.commit()

    def rollback(self):
        try:
            self.conn.rollback()
        except Exception:
            pass

    def close(self):
        if self._closed:
            return
        self._closed = True
        if self.postgres:
            conn = self.conn
            self.conn = None
            if conn is None:
                return
            try:
                # Always rollback before returning a connection to the pool so a
                # later request never inherits an unfinished transaction.
                conn.rollback()
            except Exception:
                pass
            try:
                if getattr(conn, "closed", 0):
                    self.pool.putconn(conn, close=True)
                else:
                    self.pool.putconn(conn)
            except Exception:
                try:
                    conn.close()
                except Exception:
                    pass
        else:
            self.conn.close()


def get_db():
    return CompatDB()


def ensure_column(db, table, column, definition):
    if db.postgres:
        exists = db.execute("SELECT 1 FROM information_schema.columns WHERE table_name=? AND column_name=?", (table, column)).fetchone()
        if not exists:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")
    else:
        cols = {r[1] for r in db.execute(f"PRAGMA table_info({table})").fetchall()}
        if column not in cols:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def migrate_attendance_for_lectures(db):
    """Upgrade old attendance schema to support lecture 1/2/3 per subject per day."""
    if db.postgres:
        cols = db.execute("SELECT column_name FROM information_schema.columns WHERE table_name='attendance'").fetchall()
        names = {r["column_name"] for r in cols}
        if "lecture_no" not in names:
            db.execute("ALTER TABLE attendance ADD COLUMN lecture_no INTEGER NOT NULL DEFAULT 1")
        if "lecture_time" not in names:
            db.execute("ALTER TABLE attendance ADD COLUMN lecture_time TEXT DEFAULT ''")
        # Replace the old 3-column unique constraint with the new 4-column constraint.
        constraints = db.execute("""SELECT conname FROM pg_constraint
            WHERE conrelid='attendance'::regclass AND contype='u'""").fetchall()
        for c in constraints:
            db.execute(f'ALTER TABLE attendance DROP CONSTRAINT IF EXISTS "{c["conname"]}"')
        db.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_attendance_student_subject_date_lecture ON attendance(student_id,subject_id,attendance_date,lecture_no)")
    else:
        cols = db.execute("PRAGMA table_info(attendance)").fetchall()
        names = {r[1] for r in cols}
        if "lecture_no" not in names:
            db.execute("ALTER TABLE attendance RENAME TO attendance_old")
            db.execute("""CREATE TABLE attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                attendance_date TEXT NOT NULL,
                lecture_no INTEGER NOT NULL DEFAULT 1,
                lecture_time TEXT DEFAULT '',
                status TEXT NOT NULL CHECK(status IN ('Present','Absent')),
                marked_by INTEGER,
                UNIQUE(student_id, subject_id, attendance_date, lecture_no),
                FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE,
                FOREIGN KEY(subject_id) REFERENCES subjects(id) ON DELETE CASCADE,
                FOREIGN KEY(marked_by) REFERENCES users(id) ON DELETE SET NULL
            )""")
            db.execute("""INSERT INTO attendance(id,student_id,subject_id,attendance_date,lecture_no,lecture_time,status,marked_by)
                SELECT id,student_id,subject_id,attendance_date,1,'',status,marked_by FROM attendance_old""")
            db.execute("DROP TABLE attendance_old")
        else:
            db.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_attendance_student_subject_date_lecture ON attendance(student_id,subject_id,attendance_date,lecture_no)")


def init_db():
    db = get_db()
    db.executescript("""
    CREATE TABLE IF NOT EXISTS students (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        roll_no TEXT UNIQUE NOT NULL,
        prn TEXT DEFAULT '',
        name TEXT NOT NULL,
        course TEXT NOT NULL DEFAULT 'B.Pharm',
        year TEXT NOT NULL DEFAULT '1st Year',
        division TEXT NOT NULL DEFAULT 'A',
        batch TEXT DEFAULT '',
        blocked INTEGER NOT NULL DEFAULT 0
    );

    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL,
        role TEXT NOT NULL DEFAULT 'teacher',
        full_name TEXT DEFAULT '',
        email TEXT DEFAULT '',
        mobile TEXT DEFAULT '',
        security_question TEXT DEFAULT '',
        security_answer TEXT DEFAULT '',
        approved INTEGER NOT NULL DEFAULT 1,
        student_id INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE
    );

    CREATE TABLE IF NOT EXISTS subjects (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT UNIQUE NOT NULL,
        name TEXT NOT NULL,
        year TEXT NOT NULL DEFAULT '1st Year'
    );

    CREATE TABLE IF NOT EXISTS subject_teachers (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        teacher_id INTEGER NOT NULL,
        subject_id INTEGER NOT NULL,
        UNIQUE(teacher_id, subject_id),
        FOREIGN KEY(teacher_id) REFERENCES users(id) ON DELETE CASCADE,
        FOREIGN KEY(subject_id) REFERENCES subjects(id) ON DELETE CASCADE
    );

    CREATE TABLE IF NOT EXISTS attendance (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER NOT NULL,
        subject_id INTEGER NOT NULL,
        attendance_date TEXT NOT NULL,
        lecture_no INTEGER NOT NULL DEFAULT 1,
        lecture_time TEXT DEFAULT '',
        session_type TEXT NOT NULL DEFAULT 'Lecture',
        batch TEXT DEFAULT '',
        status TEXT NOT NULL CHECK(status IN ('Present','Absent')),
        marked_by INTEGER,
        UNIQUE(student_id, subject_id, attendance_date, lecture_no),
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE,
        FOREIGN KEY(subject_id) REFERENCES subjects(id) ON DELETE CASCADE,
        FOREIGN KEY(marked_by) REFERENCES users(id) ON DELETE SET NULL
    );

    CREATE TABLE IF NOT EXISTS activity_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        action TEXT NOT NULL,
        details TEXT DEFAULT '',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS attendance_corrections (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        attendance_id INTEGER,
        student_id INTEGER NOT NULL,
        subject_id INTEGER NOT NULL,
        attendance_date TEXT NOT NULL,
        lecture_no INTEGER NOT NULL DEFAULT 1,
        old_status TEXT,
        new_status TEXT NOT NULL,
        reason TEXT NOT NULL,
        corrected_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(attendance_id) REFERENCES attendance(id) ON DELETE SET NULL,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE,
        FOREIGN KEY(subject_id) REFERENCES subjects(id) ON DELETE CASCADE,
        FOREIGN KEY(corrected_by) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS app_settings (
        key TEXT PRIMARY KEY,
        value TEXT DEFAULT ''
    );
    CREATE TABLE IF NOT EXISTS departments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT UNIQUE NOT NULL,
        name TEXT UNIQUE NOT NULL,
        active INTEGER NOT NULL DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS timetable (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        department_id INTEGER,
        year TEXT NOT NULL,
        day_of_week TEXT NOT NULL,
        lecture_no INTEGER NOT NULL,
        start_time TEXT NOT NULL,
        end_time TEXT NOT NULL,
        subject_id INTEGER,
        teacher_id INTEGER,
        room TEXT DEFAULT '',
        FOREIGN KEY(subject_id) REFERENCES subjects(id) ON DELETE SET NULL,
        FOREIGN KEY(teacher_id) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS notifications (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER NOT NULL,
        title TEXT NOT NULL,
        message TEXT NOT NULL,
        link TEXT DEFAULT '',
        is_read INTEGER NOT NULL DEFAULT 0,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS leave_applications (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER NOT NULL,
        from_date TEXT NOT NULL,
        to_date TEXT NOT NULL,
        reason TEXT NOT NULL,
        status TEXT NOT NULL DEFAULT 'Pending',
        reviewed_by INTEGER,
        review_note TEXT DEFAULT '',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE,
        FOREIGN KEY(reviewed_by) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS qr_sessions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        token TEXT UNIQUE NOT NULL,
        subject_id INTEGER NOT NULL,
        year TEXT NOT NULL,
        attendance_date TEXT NOT NULL,
        lecture_no INTEGER NOT NULL,
        lecture_time TEXT DEFAULT '',
        session_type TEXT NOT NULL DEFAULT 'Lecture',
        batch TEXT DEFAULT '',
        teacher_id INTEGER NOT NULL,
        expires_at TEXT NOT NULL,
        active INTEGER NOT NULL DEFAULT 1,
        FOREIGN KEY(subject_id) REFERENCES subjects(id) ON DELETE CASCADE,
        FOREIGN KEY(teacher_id) REFERENCES users(id) ON DELETE CASCADE
    );
    """)

    # Migrations for databases created by earlier versions.
    ensure_column(db, "users", "full_name", "TEXT DEFAULT ''")
    ensure_column(db, "users", "email", "TEXT DEFAULT ''")
    ensure_column(db, "users", "mobile", "TEXT DEFAULT ''")
    ensure_column(db, "users", "security_question", "TEXT DEFAULT ''")
    ensure_column(db, "users", "security_answer", "TEXT DEFAULT ''")
    ensure_column(db, "users", "approved", "INTEGER NOT NULL DEFAULT 1")
    ensure_column(db, "users", "student_id", "INTEGER")
    ensure_column(db, "users", "created_at", "TEXT DEFAULT CURRENT_TIMESTAMP")
    ensure_column(db, "students", "prn", "TEXT DEFAULT ''")
    ensure_column(db, "students", "department_id", "INTEGER")
    ensure_column(db, "users", "department_id", "INTEGER")
    ensure_column(db, "subjects", "department_id", "INTEGER")
    ensure_column(db, "subjects", "year", "TEXT DEFAULT '1st Year'" )
    ensure_column(db, "attendance", "marked_by", "INTEGER")
    ensure_column(db, "attendance", "academic_year", "TEXT DEFAULT ''")
    ensure_column(db, "attendance", "session_type", "TEXT NOT NULL DEFAULT 'Lecture'")
    ensure_column(db, "attendance", "batch", "TEXT DEFAULT ''")
    ensure_column(db, "attendance", "scheduled_teacher_id", "INTEGER")
    ensure_column(db, "attendance", "conducted_by_teacher_id", "INTEGER")
    ensure_column(db, "attendance", "teacher_status", "TEXT DEFAULT 'Scheduled teacher present'")
    ensure_column(db, "attendance", "start_time", "TEXT DEFAULT ''")
    ensure_column(db, "attendance", "end_time", "TEXT DEFAULT ''")
    ensure_column(db, "qr_sessions", "session_type", "TEXT NOT NULL DEFAULT 'Lecture'")
    ensure_column(db, "qr_sessions", "batch", "TEXT DEFAULT ''")
    ensure_column(db, "qr_sessions", "scheduled_teacher_id", "INTEGER")
    ensure_column(db, "qr_sessions", "conducted_by_teacher_id", "INTEGER")
    ensure_column(db, "qr_sessions", "teacher_status", "TEXT DEFAULT 'Scheduled teacher present'")
    ensure_column(db, "qr_sessions", "start_time", "TEXT DEFAULT ''")
    ensure_column(db, "qr_sessions", "end_time", "TEXT DEFAULT ''")
    ensure_column(db, "timetable", "session_type", "TEXT NOT NULL DEFAULT 'Lecture'")
    ensure_column(db, "timetable", "batch", "TEXT DEFAULT ''")
    ensure_column(db, "students", "academic_year", "TEXT DEFAULT ''")
    ensure_column(db, "students", "batch", "TEXT DEFAULT ''")
    ensure_column(db, "students", "blocked", "INTEGER NOT NULL DEFAULT 0")
    history_id_def = "SERIAL PRIMARY KEY" if db.postgres else "INTEGER PRIMARY KEY AUTOINCREMENT"
    db.execute(f"""CREATE TABLE IF NOT EXISTS student_academic_history (
        id {history_id_def},
        student_id INTEGER NOT NULL,
        academic_year TEXT NOT NULL,
        year TEXT NOT NULL,
        promoted_at TEXT DEFAULT CURRENT_TIMESTAMP,
        promoted_by INTEGER,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE,
        FOREIGN KEY(promoted_by) REFERENCES users(id) ON DELETE SET NULL,
        UNIQUE(student_id, academic_year)
    )""")
    current_ay = db.execute("SELECT value FROM app_settings WHERE key='academic_year'").fetchone()
    current_ay_value = (current_ay["value"] if current_ay and current_ay["value"] else "2026-27")
    db.execute("INSERT INTO app_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO NOTHING", ("academic_year", current_ay_value))
    db.execute("UPDATE students SET academic_year=? WHERE academic_year IS NULL OR academic_year=''", (current_ay_value,))
    db.execute("UPDATE attendance SET academic_year=? WHERE academic_year IS NULL OR academic_year=''", (current_ay_value,))
    # Multi-lecture attendance migration: old versions had one record per student/subject/day.
    # Rebuild the attendance table once so a subject can have up to 3 lectures per day.
    migrate_attendance_for_lectures(db)

    # Indexes for fast filtering/reporting when many users access the system together.
    db.execute("CREATE INDEX IF NOT EXISTS idx_students_year ON students(year)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_students_name ON students(name)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_students_batch ON students(year,batch)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_attendance_date ON attendance(attendance_date)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_attendance_subject_date ON attendance(subject_id,attendance_date)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_attendance_student_date ON attendance(student_id,attendance_date)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_attendance_academic_year ON attendance(academic_year)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_student_academic_year ON students(academic_year)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_activity_created ON activity_logs(created_at)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_notifications_user ON notifications(user_id,is_read,created_at)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_leave_status ON leave_applications(status,created_at)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_qr_token ON qr_sessions(token,active)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_timetable_day ON timetable(year,day_of_week)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_correction_date ON attendance_corrections(attendance_date)")
    # One PRN can belong to only one student when it is not blank.
    db.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_students_prn_unique ON students(prn) WHERE prn IS NOT NULL AND prn <> ''")

    if db.execute("SELECT COUNT(*) FROM departments").fetchone()[0] == 0:
        db.execute("INSERT INTO departments(code,name,active) VALUES(?,?,1)", ("PHARM", "Pharmacy"))
    default_dept = db.execute("SELECT id FROM departments ORDER BY id LIMIT 1").fetchone()[0]
    db.execute("UPDATE students SET department_id=? WHERE department_id IS NULL", (default_dept,))
    db.execute("UPDATE subjects SET department_id=? WHERE department_id IS NULL", (default_dept,))
    db.execute("UPDATE users SET department_id=? WHERE department_id IS NULL AND role IN ('teacher','hod')", (default_dept,))

    admin_username = os.environ.get("ADMIN_USERNAME", "admin").strip()
    admin_password = os.environ.get("ADMIN_PASSWORD", "")
    admin_name = os.environ.get("ADMIN_NAME", "College Administrator").strip() or "College Administrator"
    admin_email = os.environ.get("ADMIN_EMAIL", "").strip().lower()
    if db.execute("SELECT COUNT(*) FROM users WHERE role='admin'").fetchone()[0] == 0:
        if db.postgres and not admin_password:
            db.close()
            raise RuntimeError("ADMIN_PASSWORD must be set when using PostgreSQL. Set a strong admin password in Render Environment Variables.")
        admin_password = admin_password or "admin123"
        db.execute("INSERT INTO users(username,password,role,full_name,email,approved) VALUES(?,?,?,?,?,?)",
                   (admin_username, generate_password_hash(admin_password), "admin", admin_name, admin_email, 1))
    else:
        db.execute("UPDATE users SET approved=1 WHERE role='admin'")
        if admin_email:
            db.execute("UPDATE users SET email=? WHERE role='admin' AND username=?", (admin_email, admin_username))

    # Old demo teacher is intentionally removed; teachers now register.
    db.execute("DELETE FROM users WHERE role='teacher' AND username='teacher'")

    if db.execute("SELECT COUNT(*) FROM students").fetchone()[0] == 0:
        students = [
            ("01", "01", "Rahul Sharma", "B.Pharm", "1st Year", "A"),
            ("02", "02", "Priya Patel", "B.Pharm", "1st Year", "A"),
            ("03", "03", "Ankit Yadav", "B.Pharm", "1st Year", "A"),
            ("04", "04", "Sneha Gupta", "B.Pharm", "1st Year", "A"),
            ("05", "05", "Rohan Singh", "B.Pharm", "1st Year", "A"),
        ]
        db.executemany("INSERT INTO students(roll_no,prn,name,course,year,division) VALUES(?,?,?,?,?,?)", students)

    if db.execute("SELECT COUNT(*) FROM subjects").fetchone()[0] == 0:
        subjects = [
            ("BP503T", "Pharmacology-II"),
            ("BP502T", "Pharmacognosy-II"),
            ("BP501T", "Medicinal Chemistry-II"),
        ]
        db.executemany("INSERT INTO subjects(code,name,year) VALUES(?,?,?)", [(c,n,"1st Year") for c,n in subjects])

    db.commit()
    db.close()


@app.after_request
def add_security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    if request.is_secure:
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response


def validate_password(password):
    if len(password) < 8:
        return "Password must be at least 8 characters."
    if not any(c.isupper() for c in password):
        return "Password must contain at least one uppercase letter."
    if not any(c.islower() for c in password):
        return "Password must contain at least one lowercase letter."
    if not any(c.isdigit() for c in password):
        return "Password must contain at least one number."
    return None


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped


def staff_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if session.get("role") not in ("admin", "teacher", "hod"):
            flash("This page is available only to teachers and administrators.", "error")
            return redirect(url_for("student_dashboard"))
        return view(*args, **kwargs)
    return wrapped


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if session.get("role") != "admin":
            flash("Admin access required.", "error")
            return redirect(url_for("dashboard"))
        return view(*args, **kwargs)
    return wrapped


def student_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if session.get("role") != "student":
            flash("Student access required.", "error")
            return redirect(url_for("dashboard"))
        try:
            db = get_db()
            row = db.execute("SELECT blocked FROM students WHERE id=?", (session.get("student_id"),)).fetchone()
            db.close()
            if row and int(row["blocked"] or 0) == 1:
                session.clear()
                flash("You are not registered. Please contact the college administration.", "error")
                return redirect(url_for("login"))
        except Exception:
            try: db.close()
            except Exception: pass
        return view(*args, **kwargs)
    return wrapped


def teacher_or_admin_subjects(db):
    if session.get("role") == "admin":
        return db.execute("SELECT * FROM subjects ORDER BY code").fetchall()
    if session.get("role") == "hod":
        dept = db.execute("SELECT department_id FROM users WHERE id=?", (session["user_id"],)).fetchone()
        return db.execute("SELECT * FROM subjects WHERE department_id=? ORDER BY code", (dept["department_id"],)).fetchall() if dept and dept["department_id"] else []
    return db.execute("""
        SELECT s.* FROM subjects s
        JOIN subject_teachers st ON st.subject_id=s.id
        WHERE st.teacher_id=? ORDER BY s.code
    """, (session["user_id"],)).fetchall()


def can_use_subject(db, subject_id):
    if session.get("role") == "admin":
        return True
    if session.get("role") == "hod":
        row = db.execute("SELECT 1 FROM subjects s JOIN users u ON u.department_id=s.department_id WHERE s.id=? AND u.id=?", (subject_id, session["user_id"])).fetchone()
        return bool(row)
    row = db.execute("SELECT 1 FROM subject_teachers WHERE teacher_id=? AND subject_id=?",
                     (session["user_id"], subject_id)).fetchone()
    return bool(row)


def current_student(db):
    return db.execute("""
        SELECT s.* FROM students s
        JOIN users u ON u.student_id=s.id
        WHERE u.id=? AND u.role='student'
    """, (session["user_id"],)).fetchone()

def log_activity(action, details=""):
    try:
        db = get_db()
        db.execute("INSERT INTO activity_logs(user_id,action,details,created_at) VALUES(?,?,?,?)",
                   (session.get("user_id"), action, details, datetime.utcnow().isoformat()))
        db.commit(); db.close()
    except Exception:
        pass


@app.context_processor
def inject_globals():
    unread_notifications = 0
    if session.get("user_id"):
        try:
            db = get_db()
            unread_notifications = db.execute("SELECT COUNT(*) c FROM notifications WHERE user_id=? AND is_read=0", (session["user_id"],)).fetchone()["c"]
            db.close()
        except Exception:
            pass
    today_timetable=[]
    try:
        db=get_db()
        if session.get("role")=="student":
            st=current_student(db); today_timetable=get_timetable_rows(db, role="student", user_id=session.get("user_id"), student=st) if st else []
        elif session.get("role")=="teacher": today_timetable=get_timetable_rows(db, role="teacher", user_id=session.get("user_id"))
        elif session.get("role")=="admin": today_timetable=get_timetable_rows(db, role="admin", user_id=session.get("user_id"))
        db.close()
    except Exception:
        try: db.close()
        except Exception: pass
    return {
        "today_timetable": today_timetable,
        "today": date.today().isoformat(),
        "current_user": session.get("username"),
        "college_name": COLLEGE_NAME,
        "college_tagline": COLLEGE_TAGLINE,
        "logo_file": LOGO_FILE,
        "unread_notifications": unread_notifications,
    }


@app.get("/health")
def health():
    return {"status": "ok", "service": "Y.N.P. College of Pharmacy Attendance System"}


@app.route("/", methods=["GET", "POST"])
def login():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        db = get_db()
        user = db.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        if user and check_password_hash(user["password"], password):
            if user["role"] == "student":
                blocked = db.execute("SELECT blocked FROM students WHERE id=?", (user["student_id"],)).fetchone() if user["student_id"] else None
                if blocked and int(blocked["blocked"] or 0) == 1:
                    db.close()
                    flash("You are not registered. Please contact the college administration.", "error")
                    return redirect(url_for("login"))
            db.close()
            if user["approved"] == 0:
                message = "Your teacher account is waiting for admin approval." if user["role"] == "teacher" else "Your student account is waiting for approval."
                flash(message, "error")
                return redirect(url_for("login"))
            session.clear()
            session.update(user_id=user["id"], username=user["username"], role=user["role"],
                           full_name=user["full_name"] or user["username"], student_id=user["student_id"])
            if user["role"] == "student":
                return redirect(url_for("student_dashboard"))
            if user["role"] == "hod":
                return redirect(url_for("hod_dashboard"))
            return redirect(url_for("dashboard"))
        flash("Invalid username or password.", "error")
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    db0 = get_db(); departments = db0.execute("SELECT * FROM departments WHERE active=1 ORDER BY name").fetchall(); db0.close()
    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        username = request.form.get("username", "").strip()
        department_id = request.form.get("department_id", type=int)
        email = request.form.get("email", "").strip().lower()
        mobile = request.form.get("mobile", "").strip()
        security_question = request.form.get("security_question", "").strip()
        security_answer = request.form.get("security_answer", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")
        if not all([full_name, username, email, mobile, security_question, security_answer, password, confirm]):
            flash("All fields are required. Please fill every field.", "error")
            return render_template("register.html", security_questions=SECURITY_QUESTIONS, departments=departments)
        if security_question not in SECURITY_QUESTIONS:
            flash("Please select a valid security question.", "error")
            return render_template("register.html", security_questions=SECURITY_QUESTIONS, departments=departments)
        if password != confirm:
            flash("Passwords do not match.", "error")
            return render_template("register.html", security_questions=SECURITY_QUESTIONS, departments=departments)
        password_error = validate_password(password)
        if password_error:
            flash(password_error, "error")
            return render_template("register.html", security_questions=SECURITY_QUESTIONS, departments=departments)
        db = get_db()
        try:
            if db.execute("SELECT id FROM users WHERE lower(email)=lower(?)", (email,)).fetchone():
                db.close(); flash("This email is already linked to another account. Use a different email.", "error")
                return render_template("register.html", security_questions=SECURITY_QUESTIONS, departments=departments)
            if not department_id:
                department_id = db.execute("SELECT id FROM departments ORDER BY id LIMIT 1").fetchone()[0]
            db.execute("""INSERT INTO users(username,password,role,full_name,email,mobile,security_question,security_answer,approved,department_id)
                          VALUES(?,?,?,?,?,?,?,?,0,?)""",
                       (username, generate_password_hash(password), "teacher", full_name, email, mobile,
                        security_question, generate_password_hash(security_answer.casefold()), department_id))
            db.commit(); db.close()
            flash("Teacher registration submitted. Ask the college admin to approve your account and assign subjects.", "success")
            return redirect(url_for("login"))
        except Exception as exc:
            if not isinstance(exc, sqlite3.IntegrityError) and not (POSTGRES_AVAILABLE and isinstance(exc, psycopg2.IntegrityError)):
                raise
            db.close()
            flash("Username already exists. Choose another username.", "error")
    return render_template("register.html", security_questions=SECURITY_QUESTIONS, departments=departments)


@app.route("/student/register", methods=["GET", "POST"])
def student_register():
    """Create a login account for a student already added by Admin.

    Student Management is the source of truth for PRN, name, year and batch.
    Registration only creates/links the login account; it never creates a
    second student record and never asks the student to choose a batch.
    """
    if request.method == "POST":
        prn = request.form.get("prn", "").strip()
        name = " ".join(request.form.get("name", "").strip().split())
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower()
        mobile = request.form.get("mobile", "").strip()
        security_question = request.form.get("security_question", "").strip()
        security_answer = request.form.get("security_answer", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")
        year = request.form.get("year", "").strip()
        context = dict(
            year_options=YEAR_OPTIONS,
            selected_year=year,
            security_questions=SECURITY_QUESTIONS,
            departments=[],
        )

        if not all([prn, name, email, username, mobile, security_question,
                    security_answer, password, confirm, year]):
            flash("All fields are required. Please fill every field and select your year.", "error")
            return render_template("student_register.html", **context)
        if year not in YEAR_OPTIONS:
            flash("Please select a valid year.", "error")
            return render_template("student_register.html", **context)
        if security_question not in SECURITY_QUESTIONS:
            flash("Please select a valid security question.", "error")
            return render_template("student_register.html", **context)
        if password != confirm:
            flash("Passwords do not match.", "error")
            return render_template("student_register.html", **context)
        password_error = validate_password(password)
        if password_error:
            flash(password_error, "error")
            return render_template("student_register.html", **context)

        db = get_db()
        try:
            # The Admin-created Student Management record is required.
            student = db.execute(
                "SELECT * FROM students WHERE prn=? LIMIT 1", (prn,)
            ).fetchone()
            if not student:
                db.close()
                flash("This PRN is not in Student Management. Please ask the admin to add the student first.", "error")
                return render_template("student_register.html", **context)

            if int(student["blocked"] or 0) == 1:
                db.close()
                flash("You are not registered. Please contact the college administration.", "error")
                return render_template("student_register.html", **context)

            # One student record can have only one student login account.
            linked = db.execute(
                "SELECT id FROM users WHERE student_id=? LIMIT 1", (student["id"],)
            ).fetchone()
            if linked:
                db.close()
                flash("An account for this PRN already exists. Please log in instead.", "error")
                return render_template("student_register.html", **context)

            # Registration must match the Admin-created student's identity.
            existing_name = " ".join(str(student["name"] or "").strip().split())
            if existing_name.casefold() != name.casefold():
                db.close()
                flash("The name does not match the student record. Please enter the name exactly as added by the admin.", "error")
                return render_template("student_register.html", **context)
            if str(student["year"] or "").strip() != year:
                db.close()
                flash("The selected year does not match the student record. Please select the year assigned by the admin.", "error")
                return render_template("student_register.html", **context)

            if db.execute("SELECT id FROM users WHERE lower(email)=lower(?)", (email,)).fetchone():
                db.close()
                flash("This email is already linked to another account. Use a different email.", "error")
                return render_template("student_register.html", **context)
            if db.execute("SELECT id FROM users WHERE lower(username)=lower(?)", (username,)).fetchone():
                db.close()
                flash("This username is already in use. Choose another username.", "error")
                return render_template("student_register.html", **context)

            # Reuse the student's existing internal department relationship.
            department_id = student["department_id"] if student["department_id"] is not None else None
            db.execute(
                """INSERT INTO users(username,password,role,full_name,email,mobile,
                   security_question,security_answer,approved,student_id,department_id)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (username, generate_password_hash(password), "student", student["name"],
                 email, mobile, security_question,
                 generate_password_hash(security_answer.casefold()), 1,
                 student["id"], department_id),
            )
            db.commit()
            db.close()
            flash("Student account created successfully. You can now log in.", "success")
            return redirect(url_for("login"))
        except Exception as exc:
            try:
                db.close()
            except Exception:
                pass
            if not isinstance(exc, sqlite3.IntegrityError) and not (POSTGRES_AVAILABLE and isinstance(exc, psycopg2.IntegrityError)):
                raise
            flash("Username, email or account details already exist. Please use different details.", "error")

    return render_template(
        "student_register.html",
        year_options=YEAR_OPTIONS,
        selected_year="",
        security_questions=SECURITY_QUESTIONS,
        departments=[],
    )


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    # Do not carry login-page/block messages into the password-recovery page.
    # They should be shown on the login page only.
    if request.method == "GET":
        session.pop("_flashes", None)
    if request.method == "POST":
        step = request.form.get("step", "1")
        if step == "2":
            user_id = session.get("password_reset_user_id")
            if not user_id:
                flash("Recovery session expired. Please start again.", "error")
                return redirect(url_for("forgot_password"))
            answer = request.form.get("security_answer", "").strip().casefold()
            attempts = int(session.get("password_reset_attempts", 0)) + 1
            session["password_reset_attempts"] = attempts
            db = get_db()
            user = db.execute("SELECT id,security_answer,approved FROM users WHERE id=?", (user_id,)).fetchone()
            db.close()
            if attempts > 5:
                session.pop("password_reset_user_id", None); session.pop("password_reset_attempts", None)
                flash("Too many incorrect attempts. Please start password recovery again.", "error")
                return redirect(url_for("forgot_password"))
            if user and user["approved"] and user["security_answer"] and check_password_hash(user["security_answer"], answer):
                session["password_reset_verified"] = True
                return redirect(url_for("reset_password"))
            flash(f"Incorrect security answer. Attempts remaining: {max(0, 5-attempts)}", "error")
            db = get_db(); row = db.execute("SELECT security_question FROM users WHERE id=?", (user_id,)).fetchone(); db.close()
            return render_template("forgot_password.html", step=2, security_question=row["security_question"] if row else "")

        account_type = request.form.get("account_type", "student")
        identifier = request.form.get("identifier", "").strip()
        mobile = request.form.get("mobile", "").strip()
        if not identifier or not mobile:
            flash("Please enter all required details.", "error")
            return render_template("forgot_password.html", step=1)
        db = get_db()
        if account_type == "teacher":
            user = db.execute("SELECT id,security_question,approved FROM users WHERE lower(username)=lower(?) AND mobile=? AND role='teacher' LIMIT 1", (identifier, mobile)).fetchone()
        else:
            user = db.execute("""SELECT u.id,u.security_question,u.approved FROM users u
                                JOIN students s ON s.id=u.student_id
                                WHERE s.prn=? AND u.mobile=? AND u.role='student' LIMIT 1""", (identifier, mobile)).fetchone()
        db.close()
        session.pop("password_reset_verified", None)
        session.pop("password_reset_user_id", None)
        session.pop("password_reset_attempts", None)
        if not user or not user["approved"] or not user["security_question"]:
            flash("The details do not match an approved account, or security recovery is not set up.", "error")
            return render_template("forgot_password.html", step=1)
        session["password_reset_user_id"] = user["id"]
        session["password_reset_attempts"] = 0
        return render_template("forgot_password.html", step=2, security_question=user["security_question"])
    return render_template("forgot_password.html", step=1)


@app.route("/reset-password", methods=["GET", "POST"])
def reset_password():
    if not session.get("password_reset_verified") or not session.get("password_reset_user_id"):
        flash("Please verify your recovery details first.", "error")
        return redirect(url_for("forgot_password"))
    user_id = session["password_reset_user_id"]
    if request.method == "POST":
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")
        password_error = validate_password(password)
        if password_error:
            flash(password_error, "error")
            return render_template("reset_password.html")
        if password != confirm:
            flash("Passwords do not match.", "error")
            return render_template("reset_password.html")
        db = get_db()
        db.execute("UPDATE users SET password=? WHERE id=?", (generate_password_hash(password), user_id))
        db.commit(); db.close()
        session.pop("password_reset_verified", None); session.pop("password_reset_user_id", None); session.pop("password_reset_attempts", None)
        flash("Password changed successfully. You can now log in with your new password.", "success")
        return redirect(url_for("login"))
    return render_template("reset_password.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


def get_timetable_rows(db, role=None, user_id=None, student=None, day_name=None):
    day_name = day_name or date.today().strftime("%A")
    where=["t.day_of_week=?"]; params=[day_name]
    if student is not None:
        where.append("t.year=?"); params.append(student["year"])
        where.append("(COALESCE(t.session_type,'Lecture')='Lecture' OR COALESCE(t.batch,'')=?)"); params.append(student["batch"] or "")
    elif role == "teacher":
        where.append("t.teacher_id=?"); params.append(user_id)
    elif role == "hod":
        where.append("t.department_id=(SELECT department_id FROM users WHERE id=? )"); params.append(user_id)
    return db.execute(f"""SELECT t.*, COALESCE(t.session_type,'Lecture') session_type, COALESCE(t.batch,'') batch,
        COALESCE(sub.code,'') subject_code, COALESCE(sub.name,'') subject_name,
        COALESCE(u.full_name,u.username,'') teacher_name
        FROM timetable t LEFT JOIN subjects sub ON sub.id=t.subject_id LEFT JOIN users u ON u.id=t.teacher_id
        WHERE {' AND '.join(where)} ORDER BY t.start_time,t.lecture_no,t.id""", tuple(params)).fetchall()


def timetable_candidates(db, year, session_type, batch, subject_id, day_name):
    return db.execute("""SELECT t.*, COALESCE(u.full_name,u.username,'') teacher_name,
        COALESCE(sub.code,'') subject_code, COALESCE(sub.name,'') subject_name
        FROM timetable t LEFT JOIN users u ON u.id=t.teacher_id LEFT JOIN subjects sub ON sub.id=t.subject_id
        WHERE t.year=? AND t.day_of_week=? AND t.subject_id=?
          AND COALESCE(t.session_type,'Lecture')=?
          AND (COALESCE(t.batch,'')=? OR ?='Lecture')
        ORDER BY t.start_time,t.lecture_no,t.id""",
        (year,day_name,subject_id,session_type,batch if session_type=='Practical' else '',session_type)).fetchall()


def current_academic_year_value(db):
    row=db.execute("SELECT value FROM app_settings WHERE key='academic_year'").fetchone()
    return row['value'] if row and row['value'] else '2026-27'


def all_teacher_rows(db):
    return db.execute("SELECT id,COALESCE(full_name,username) name FROM users WHERE role='teacher' AND approved=1 ORDER BY name").fetchall()


@app.route("/timetable", methods=["GET", "POST"])
@login_required
def timetable():
    db=get_db(); role=session.get("role")
    days=['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday']
    if request.method=='POST':
        if role!='admin':
            db.close(); flash('Only administrators can manage the timetable.','error'); return redirect(url_for('timetable'))
        year=request.form.get('year','').strip(); day=request.form.get('day_of_week','').strip()
        start=request.form.get('start_time','').strip(); end=request.form.get('end_time','').strip()
        subject_id=request.form.get('subject_id',type=int); teacher_id=request.form.get('teacher_id',type=int)
        session_type=request.form.get('session_type','Lecture').strip(); batch=request.form.get('batch','').strip(); room=request.form.get('room','').strip()
        if year not in YEAR_OPTIONS or day not in days or not start or not end or not subject_id or not teacher_id or session_type not in ('Lecture','Practical'):
            flash('Please fill all timetable fields correctly.','error')
        elif session_type=='Practical' and batch not in ('Batch A','Batch B','Batch C','Batch D'):
            flash('Select a valid practical batch.','error')
        else:
            dept_row=db.execute('SELECT id FROM departments ORDER BY id LIMIT 1').fetchone(); dept=dept_row['id'] if dept_row else None
            db.execute("INSERT INTO timetable(department_id,year,day_of_week,lecture_no,start_time,end_time,subject_id,teacher_id,room,session_type,batch) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                       (dept,year,day,1,start,end,subject_id,teacher_id,room,session_type,batch if session_type=='Practical' else ''))
            db.commit(); flash('Timetable entry added successfully.','success')
    selected_year=request.args.get('year','').strip(); selected_day=request.args.get('day','').strip()
    if selected_day not in days: selected_day=date.today().strftime('%A')
    where=[]; params=[]
    if role=='teacher': where.append('t.teacher_id=?'); params.append(session['user_id'])
    elif role=='student':
        st=current_student(db); selected_year=st['year'] if st else selected_year
        where += ['t.year=?',"(COALESCE(t.session_type,'Lecture')='Lecture' OR COALESCE(t.batch,'')=?)"]; params += [selected_year, st['batch'] if st else '']
    elif role=='hod': where.append('t.department_id=(SELECT department_id FROM users WHERE id=?)'); params.append(session['user_id'])
    if selected_year and role in ('admin','hod'): where.append('t.year=?'); params.append(selected_year)
    where.append('t.day_of_week=?'); params.append(selected_day)
    rows=db.execute(f"""SELECT t.*,COALESCE(t.session_type,'Lecture') session_type,COALESCE(t.batch,'') batch,
        COALESCE(sub.code,'') subject_code,COALESCE(sub.name,'') subject_name,COALESCE(u.full_name,u.username,'') teacher_name
        FROM timetable t LEFT JOIN subjects sub ON sub.id=t.subject_id LEFT JOIN users u ON u.id=t.teacher_id
        WHERE {' AND '.join(where)} ORDER BY t.year,t.start_time,t.lecture_no,t.id""",tuple(params)).fetchall()
    subjects=db.execute('SELECT id,code,name,year FROM subjects ORDER BY year,code').fetchall(); teachers=all_teacher_rows(db)
    db.close()
    return render_template('timetable.html',rows=rows,subjects=subjects,teachers=teachers,years=YEAR_OPTIONS,selected_year=selected_year,selected_day=selected_day,days=days,batch_options=['Batch A','Batch B','Batch C','Batch D'])


@app.route('/timetable/export.pdf')
@login_required
def export_timetable_pdf():
    """Export the timetable currently visible to the logged-in user as a professional PDF."""
    db=get_db(); role=session.get('role')
    days=['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday']
    selected_day=request.args.get('day','').strip()
    if selected_day not in days:
        selected_day=date.today().strftime('%A')
        if selected_day not in days:
            selected_day='Monday'
    selected_year=request.args.get('year','').strip()

    where=[]; params=[]
    if role=='teacher':
        where.append('t.teacher_id=?'); params.append(session['user_id'])
    elif role=='student':
        st=current_student(db)
        if not st:
            db.close(); return Response('Student record not found.',status=404)
        selected_year=st['year']
        where += ['t.year=?', "(COALESCE(t.session_type,'Lecture')='Lecture' OR COALESCE(t.batch,'')=?)"]
        params += [selected_year, st['batch'] or '']
    elif role=='hod':
        where.append('t.department_id=(SELECT department_id FROM users WHERE id=?)'); params.append(session['user_id'])
    elif role=='admin':
        if selected_year and selected_year in YEAR_OPTIONS:
            where.append('t.year=?'); params.append(selected_year)
    else:
        db.close(); return Response('Not authorized.',status=403)

    where.append('t.day_of_week=?'); params.append(selected_day)
    rows=db.execute(f"""SELECT t.*, COALESCE(t.session_type,'Lecture') session_type,
        COALESCE(t.batch,'') batch, COALESCE(sub.code,'') subject_code,
        COALESCE(sub.name,'') subject_name, COALESCE(u.full_name,u.username,'') teacher_name
        FROM timetable t LEFT JOIN subjects sub ON sub.id=t.subject_id
        LEFT JOIN users u ON u.id=t.teacher_id
        WHERE {' AND '.join(where)}
        ORDER BY t.year,t.start_time,t.lecture_no,t.id""",tuple(params)).fetchall()
    db.close()

    buffer=io.BytesIO()
    doc=SimpleDocTemplate(buffer,pagesize=landscape(A4),rightMargin=10*mm,leftMargin=10*mm,topMargin=9*mm,bottomMargin=9*mm)
    styles=getSampleStyleSheet()
    title=ParagraphStyle('TTCollegeTitle',parent=styles['Title'],alignment=TA_CENTER,fontSize=19,leading=22,spaceAfter=1)
    sub=ParagraphStyle('TTSub',parent=styles['Normal'],alignment=TA_CENTER,fontSize=10.5,leading=13,spaceAfter=1)
    meta=ParagraphStyle('TTMeta',parent=styles['Normal'],alignment=TA_CENTER,fontSize=9.5,leading=12)
    cell=ParagraphStyle('TTCell',parent=styles['Normal'],fontSize=7.8,leading=9.2,alignment=TA_LEFT)
    head=ParagraphStyle('TTHead',parent=cell,fontName='Helvetica-Bold',alignment=TA_CENTER)

    story=[]
    logo_path=BASE_DIR/'static'/LOGO_FILE
    logo=Image(str(logo_path),width=23*mm,height=23*mm) if logo_path.exists() else ''
    class_text=selected_year if selected_year else 'All Years'
    header_text=[Paragraph(COLLEGE_NAME,title),
                 Paragraph('Academic Timetable',sub),
                 Paragraph('Affiliated by Dr. Babasaheb Ambedkar Technological University, Lonere, Raigad',meta),
                 Paragraph(f'<b>Day:</b> {selected_day} &nbsp;&nbsp; <b>Year:</b> {class_text}',meta)]
    hdr=Table([[logo,header_text]],colWidths=[28*mm,238*mm])
    hdr.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),0)]))
    story.append(hdr); story.append(Spacer(1,4*mm))

    data=[[Paragraph('Year',head),Paragraph('Time',head),Paragraph('Session',head),Paragraph('Batch',head),Paragraph('Subject',head),Paragraph('Teacher',head),Paragraph('Room / Lab',head)]]
    for r in rows:
        subject=f"{r['subject_code']} — {r['subject_name']}" if r['subject_code'] else (r['subject_name'] or '—')
        data.append([Paragraph(str(r['year'] or '—'),cell),Paragraph(f"<b>{r['start_time']}</b> – <b>{r['end_time']}</b>",cell),Paragraph(str(r['session_type'] or 'Lecture'),cell),Paragraph(r['batch'] or 'All',cell),Paragraph(subject,cell),Paragraph(r['teacher_name'] or '—',cell),Paragraph(r['room'] or '—',cell)])
    if not rows:
        data.append([Paragraph('No timetable entries found for the selected filter.',cell),'','','','','',''])
    table=Table(data,repeatRows=1,colWidths=[24*mm,34*mm,27*mm,24*mm,67*mm,50*mm,38*mm])
    table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0b3974')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#b8c7d9')),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),
        ('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),
        ('ALIGN',(1,1),(3,-1),'CENTER'),
    ]))
    story.append(table)
    story.append(Spacer(1,4*mm))
    generated=datetime.now(ZoneInfo('Asia/Kolkata')).strftime('%d-%m-%Y %I:%M %p')
    story.append(Paragraph(f'Generated on: {generated} IST &nbsp;&nbsp; | &nbsp;&nbsp; Y.N.P. College of Pharmacy',meta))
    doc.build(story)
    buffer.seek(0)
    filename=f"timetable_{(selected_year or 'all').replace(' ','_')}_{selected_day.lower()}.pdf"
    return Response(buffer.getvalue(),mimetype='application/pdf',headers={'Content-Disposition':f'attachment; filename={filename}'})


@app.post('/timetable/delete/<int:timetable_id>')
@login_required
@admin_required
def delete_timetable(timetable_id):
    db=get_db(); db.execute('DELETE FROM timetable WHERE id=?',(timetable_id,)); db.commit(); db.close(); flash('Timetable entry deleted.','success'); return redirect(url_for('timetable'))

@app.route("/dashboard")
@login_required
def dashboard():
    if session.get("role") == "student":
        return redirect(url_for("student_dashboard"))
    if session.get("role") == "hod":
        return redirect(url_for("hod_dashboard"))
    db = get_db()
    today = date.today().isoformat()
    total = db.execute("SELECT COUNT(*) c FROM students").fetchone()["c"]
    pending = db.execute("SELECT COUNT(*) c FROM users WHERE role='teacher' AND approved=0").fetchone()["c"] if session.get("role") == "admin" else 0

    subject_filter = "" if session.get("role") == "admin" else " AND a.subject_id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?)"
    base_params = [] if session.get("role") == "admin" else [session["user_id"]]
    today_rows = db.execute(f"""SELECT a.student_id,a.status,a.lecture_no,s.year
        FROM attendance a JOIN students s ON s.id=a.student_id
        WHERE a.attendance_date=?{subject_filter}""", [today] + base_params).fetchall()
    present_marks = sum(1 for r in today_rows if r["status"] == "Present")
    total_marks = len(today_rows)
    absent_marks = total_marks - present_marks
    today_pct = round(100.0 * present_marks / total_marks, 1) if total_marks else 0

    year_stats=[]
    for y in YEAR_OPTIONS:
        yr_students = db.execute("SELECT COUNT(*) c FROM students WHERE year=?", (y,)).fetchone()["c"]
        rows = [r for r in today_rows if r["year"] == y]
        p = sum(1 for r in rows if r["status"] == "Present")
        a = sum(1 for r in rows if r["status"] == "Absent")
        marked_students = {r["student_id"] for r in rows}
        # A student is counted as present if present for every marked lecture; otherwise absent.
        per_student = {}
        for r in rows:
            per_student.setdefault(r["student_id"], []).append(r["status"])
        present_students = sum(1 for sid, sts in per_student.items() if sts and all(x == "Present" for x in sts))
        absent_students = sum(1 for sid, sts in per_student.items() if any(x == "Absent" for x in sts))
        pct = round(100.0 * p / (p+a), 1) if (p+a) else 0
        year_stats.append({"year":y,"total":yr_students,"present":present_students,"absent":absent_students,"pct":pct,"marked":len(marked_students),"present_marks":p,"absent_marks":a})

    threshold_row = db.execute("SELECT value FROM app_settings WHERE key='attendance_threshold'").fetchone()
    try: threshold = float(threshold_row["value"] if threshold_row and threshold_row["value"] else 75)
    except Exception: threshold = 75
    low = db.execute("""SELECT s.prn,s.name,s.course,s.year,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM students s LEFT JOIN attendance a ON a.student_id=s.id
        GROUP BY s.id HAVING (CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END) < ?
        ORDER BY pct ASC, s.name LIMIT 10""", (threshold,)).fetchall()
    # Overall attendance uses lecture marks, so multiple lectures are counted correctly.
    avg = db.execute("""SELECT COALESCE(AVG(pct),0) avg_pct FROM (
        SELECT s.id, CASE WHEN COUNT(a.id)=0 THEN 0 ELSE 100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id) END pct
        FROM students s LEFT JOIN attendance a ON a.student_id=s.id GROUP BY s.id)""").fetchone()["avg_pct"]
    if session.get("role") == "admin":
        subject_count = db.execute("SELECT COUNT(*) c FROM subjects").fetchone()["c"]
        teacher_count = db.execute("SELECT COUNT(*) c FROM users WHERE role='teacher' AND approved=1").fetchone()["c"]
    else:
        subject_count = db.execute("SELECT COUNT(*) c FROM subject_teachers WHERE teacher_id=?", (session["user_id"],)).fetchone()["c"]
        teacher_count = 0

    recent_attendance = db.execute("""SELECT a.attendance_date, sub.code, sub.name subject_name, s.year,
        COUNT(a.id) total,
        SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END) present,
        SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END) absent
        FROM attendance a JOIN subjects sub ON sub.id=a.subject_id JOIN students s ON s.id=a.student_id
        WHERE 1=1
        """ + (" AND a.subject_id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?) " if session.get('role')=='teacher' else "") + """GROUP BY a.attendance_date, sub.id, sub.code, sub.name, s.year
        ORDER BY a.attendance_date DESC LIMIT 6""", ([session["user_id"]] if session.get('role')=='teacher' else [])).fetchall()
    db.close()
    return render_template("dashboard.html", total=total, present=present_marks, absent=absent_marks,
        avg=round(avg or 0,1), low=low, pending=pending, year_stats=year_stats, today=today, today_pct=today_pct, today_marks=total_marks, threshold=threshold, subject_count=subject_count, teacher_count=teacher_count, recent_attendance=recent_attendance)


@app.route("/student/dashboard")
@student_required
def student_dashboard():
    db = get_db()
    student = current_student(db)
    if not student:
        db.close()
        session.clear()
        flash("Student account is not linked correctly. Contact the administrator.", "error")
        return redirect(url_for("login"))

    today = date.today().isoformat()
    selected_date = request.args.get("date", "").strip()
    selected_subject = request.args.get("subject", "").strip()
    timetable_days = ['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday']
    selected_timetable_day = request.args.get("timetable_day", "").strip()
    if selected_timetable_day not in timetable_days:
        selected_timetable_day = date.today().strftime('%A')
        if selected_timetable_day not in timetable_days:
            selected_timetable_day = 'Monday'

    today_rows = db.execute("""SELECT sub.code, sub.name subject_name, a.status,
        COALESCE(a.lecture_no,1) lecture_no, COALESCE(a.lecture_time,'') lecture_time
        FROM subjects sub LEFT JOIN attendance a
        ON a.subject_id=sub.id AND a.student_id=? AND a.attendance_date=?
        WHERE sub.department_id=? AND sub.year=?
        ORDER BY sub.code, a.lecture_no""", (student["id"], today, student["department_id"], student["year"])).fetchall()

    subjects = db.execute("SELECT id, code, name FROM subjects WHERE department_id=? AND year=? ORDER BY code", (student["department_id"], student["year"])).fetchall()

    history_sql = """SELECT a.attendance_date, sub.code, sub.name subject_name,
        a.status, COALESCE(a.lecture_no,1) lecture_no, COALESCE(a.lecture_time,'') lecture_time
        FROM attendance a JOIN subjects sub ON sub.id=a.subject_id
        WHERE a.student_id=?"""
    history_params = [student["id"]]
    if selected_date:
        history_sql += " AND a.attendance_date=?"
        history_params.append(selected_date)
    if selected_subject:
        history_sql += " AND a.subject_id=?"
        history_params.append(selected_subject)
    history_sql += " ORDER BY a.attendance_date DESC, sub.code, a.lecture_no"
    history = db.execute(history_sql, tuple(history_params)).fetchall()

    summary = db.execute("""SELECT sub.code, sub.name subject_name,
        COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM subjects sub LEFT JOIN attendance a ON a.subject_id=sub.id AND a.student_id=?
        WHERE sub.department_id=? AND sub.year=? GROUP BY sub.id ORDER BY sub.code""", (student["id"], student["department_id"], student["year"])).fetchall()

    total = db.execute("SELECT COUNT(*) FROM attendance WHERE student_id=?", (student["id"],)).fetchone()[0]
    present = db.execute("SELECT COUNT(*) FROM attendance WHERE student_id=? AND status='Present'", (student["id"],)).fetchone()[0]
    overall = round((present / total * 100), 1) if total else 0

    # Student dashboard timetable is strictly limited to the logged-in
    # student's current year and practical batch. Lectures are shown for the
    # year, while practicals are shown only for the student's batch.
    my_timetable = get_timetable_rows(
        db, role='student', user_id=session.get('user_id'),
        student=student, day_name=selected_timetable_day
    )
    db.close()
    return render_template("student_dashboard.html", student=student, today_rows=today_rows,
                           summary=summary, overall=overall, history=history, subjects=subjects,
                           selected_date=selected_date, selected_subject=selected_subject, today=today,
                           my_timetable=my_timetable, timetable_days=timetable_days,
                           selected_timetable_day=selected_timetable_day)


@app.route("/student/calendar")
@student_required
def student_calendar():
    db=get_db(); student=current_student(db)
    month=request.args.get("month",date.today().strftime("%Y-%m"))
    if not re.match(r"^\d{4}-\d{2}$",month): month=date.today().strftime("%Y-%m")
    start=f"{month}-01"
    y,m=map(int,month.split("-")); end=f"{y+1:04d}-01-01" if m==12 else f"{y:04d}-{m+1:02d}-01"
    rows=db.execute("""SELECT a.attendance_date, SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END) present, COUNT(a.id) total
        FROM attendance a WHERE a.student_id=? AND a.attendance_date>=? AND a.attendance_date<? GROUP BY a.attendance_date ORDER BY a.attendance_date""",(student["id"],start,end)).fetchall()
    db.close(); return render_template("student_calendar.html",student=student,month=month,days=rows)

@app.route("/students")
@login_required
@admin_required
def students():
    q = request.args.get("q", "").strip()
    selected_year = request.args.get("year", "").strip()
    db = get_db()

    # Four year-wise student cards. Counts come directly from the students table.
    year_counts = {}
    for year in YEAR_OPTIONS:
        year_counts[year] = db.execute(
            "SELECT COUNT(*) FROM students WHERE year=?", (year,)
        ).fetchone()[0]

    conditions = []
    params = []
    if selected_year in YEAR_OPTIONS:
        conditions.append("s.year=?")
        params.append(selected_year)
    if q:
        conditions.append("(s.prn LIKE ? OR s.name LIKE ? OR s.course LIKE ?)")
        params.extend([f"%{q}%", f"%{q}%", f"%{q}%"])

    where = (" WHERE " + " AND ".join(conditions)) if conditions else ""
    rows = db.execute(f"""SELECT s.*, u.username student_username
        FROM students s LEFT JOIN users u ON u.student_id=s.id AND u.role='student'
        {where}
        ORDER BY s.prn, s.name""", tuple(params)).fetchall()

    # Arrange students by complete PRN in ascending numerical order.
    # Numeric PRNs are compared as numbers, so ...3009, ...3010, ...3011
    # comes before ...3020, ...3024, etc. Non-numeric PRNs are kept after
    # numeric PRNs and sorted alphabetically.
    def _prn_sort_key(row):
        value = str(row["prn"] or "").strip()
        if value.isdigit():
            return (0, int(value), "")
        return (1, 0, value.lower())

    rows = sorted(rows, key=_prn_sort_key)
    db.close()
    return render_template("students.html", students=rows, q=q, selected_year=selected_year,
                           year_options=YEAR_OPTIONS, year_counts=year_counts)


@app.route("/students/add", methods=["GET", "POST"])
@login_required
@admin_required
def add_student():
    db = get_db()
    departments = db.execute("SELECT * FROM departments WHERE active=1 ORDER BY name").fetchall()
    if request.method == "POST":
        data = [request.form.get(k, "").strip() for k in ("prn", "name", "course", "year")]
        batch = request.form.get("batch", "").strip()
        dept_id = departments[0]["id"] if departments else None
        if not data[0] or not data[1] or data[3] not in YEAR_OPTIONS or not dept_id:
            db.close(); flash("PRN, name and year are required.", "error")
            return render_template("student_form.html", student=None, title="Add Student", year_options=YEAR_OPTIONS, batch_options=["Batch A","Batch B","Batch C","Batch D"])
        try:
            db.execute("INSERT INTO students(roll_no,prn,name,course,year,division,department_id,batch) VALUES(?,?,?,?,?,?,?,?)", (data[0], data[0], data[1], data[2] or "B.Pharm", data[3], "A", dept_id, batch))
            db.commit(); flash("Student added successfully.", "success")
        except Exception as exc:
            if not isinstance(exc, sqlite3.IntegrityError) and not (POSTGRES_AVAILABLE and isinstance(exc, psycopg2.IntegrityError)): raise
            flash("PRN already exists.", "error")
        db.close(); return redirect(url_for("students"))
    db.close(); return render_template("student_form.html", student=None, title="Add Student", year_options=YEAR_OPTIONS, batch_options=["Batch A","Batch B","Batch C","Batch D"])


@app.route("/students/edit/<int:student_id>", methods=["GET", "POST"])
@login_required
@admin_required
def edit_student(student_id):
    db = get_db(); student = db.execute("SELECT * FROM students WHERE id=?", (student_id,)).fetchone()
    if not student:
        db.close(); flash("Student not found.", "error"); return redirect(url_for("students"))
    if request.method == "POST":
        data = [request.form.get(k, "").strip() for k in ("prn", "name", "course", "year")]
        batch = request.form.get("batch", "").strip()
        try:
            dept_id = student["department_id"]
            if not dept_id:
                dept_row = db.execute("SELECT id FROM departments ORDER BY id LIMIT 1").fetchone()
                dept_id = dept_row["id"] if dept_row else None
            if not data[0] or not data[1] or data[3] not in YEAR_OPTIONS:
                raise ValueError("PRN, name and valid year are required.")
            db.execute("UPDATE students SET roll_no=?,prn=?,name=?,course=?,year=?,department_id=?,batch=? WHERE id=?", (data[0], data[0], data[1], data[2] or "B.Pharm", data[3], dept_id, batch, student_id))
            db.execute("UPDATE users SET full_name=? WHERE student_id=? AND role='student'", (data[1], student_id))
            db.commit(); flash("Student updated successfully.", "success"); db.close(); return redirect(url_for("students"))
        except ValueError as exc:
            db.close(); flash(str(exc), "error")
        except Exception as exc:
            if not isinstance(exc, sqlite3.IntegrityError) and not (POSTGRES_AVAILABLE and isinstance(exc, psycopg2.IntegrityError)):
                raise
            flash("PRN already exists.", "error")
    departments=db.execute("SELECT * FROM departments WHERE active=1 ORDER BY name").fetchall(); db.close(); return render_template("student_form.html", student=student, title="Edit Student", year_options=YEAR_OPTIONS, batch_options=["Batch A","Batch B","Batch C","Batch D"])


@app.route("/students/delete/<int:student_id>", methods=["POST"])
@login_required
@admin_required
def delete_student(student_id):
    db = get_db(); db.execute("DELETE FROM students WHERE id=?", (student_id,)); db.commit(); db.close()
    flash("Student and linked attendance/account were deleted.", "success"); return redirect(url_for("students"))


@app.route("/students/<int:student_id>/toggle-block", methods=["POST"])
@login_required
@admin_required
def toggle_block_student(student_id):
    db = get_db()
    student = db.execute("SELECT id, name, blocked FROM students WHERE id=?", (student_id,)).fetchone()
    if not student:
        db.close()
        flash("Student not found.", "error")
        return redirect(url_for("students"))
    new_status = 0 if int(student["blocked"] or 0) else 1
    db.execute("UPDATE students SET blocked=? WHERE id=?", (new_status, student_id))
    db.commit()
    db.close()
    if new_status:
        flash(f"{student['name']} has been blocked. Student login and registration are disabled.", "success")
    else:
        flash(f"{student['name']} has been unblocked. Student can log in and register again.", "success")
    return redirect(url_for("students"))


@app.route("/teachers")
@login_required
@admin_required
def teachers():
    db = get_db()
    rows = db.execute("SELECT * FROM users WHERE role='teacher' ORDER BY approved ASC, full_name, username").fetchall()
    subjects = db.execute("SELECT * FROM subjects ORDER BY year, code").fetchall()
    assignments = db.execute("SELECT teacher_id, subject_id FROM subject_teachers").fetchall()
    assigned = {(r["teacher_id"], r["subject_id"]) for r in assignments}
    db.close()
    return render_template("teachers.html", teachers=rows, subjects=subjects, assigned=assigned, years=YEAR_OPTIONS)


@app.route("/teachers/<int:teacher_id>/approve", methods=["POST"])
@login_required
@admin_required
def approve_teacher(teacher_id):
    db = get_db(); db.execute("UPDATE users SET approved=1 WHERE id=? AND role='teacher'", (teacher_id,)); db.commit(); db.close()
    flash("Teacher account approved.", "success"); return redirect(url_for("teachers"))


@app.route("/teachers/<int:teacher_id>/toggle", methods=["POST"])
@login_required
@admin_required
def toggle_teacher(teacher_id):
    db = get_db(); db.execute("UPDATE users SET approved=CASE approved WHEN 1 THEN 0 ELSE 1 END WHERE id=? AND role='teacher'", (teacher_id,)); db.commit(); db.close()
    flash("Teacher account status updated.", "success"); return redirect(url_for("teachers"))


@app.route("/teachers/<int:teacher_id>/delete", methods=["POST"])
@login_required
@admin_required
def delete_teacher(teacher_id):
    db = get_db()
    teacher = db.execute("SELECT full_name,username FROM users WHERE id=? AND role='teacher'", (teacher_id,)).fetchone()
    if not teacher:
        db.close(); flash("Teacher not found.", "error"); return redirect(url_for("teachers"))
    db.execute("DELETE FROM users WHERE id=? AND role='teacher'", (teacher_id,))
    db.commit(); db.close()
    flash(f"Teacher {teacher['full_name'] or teacher['username']} removed successfully.", "success")
    return redirect(url_for("teachers"))


@app.route("/teachers/<int:teacher_id>/assign", methods=["POST"])
@login_required
@admin_required
def assign_subjects(teacher_id):
    db = get_db()
    db.execute("DELETE FROM subject_teachers WHERE teacher_id=?", (teacher_id,))
    subject_ids = request.form.getlist("subject_ids")
    valid = db.execute("SELECT id FROM subjects").fetchall()
    valid_ids = {str(r["id"]) for r in valid}
    for sid in subject_ids:
        if sid in valid_ids:
            db.execute("INSERT OR IGNORE INTO subject_teachers(teacher_id,subject_id) VALUES(?,?)", (teacher_id, int(sid)))
    db.commit(); db.close()
    flash("Teacher subject assignments saved year-wise.", "success"); return redirect(url_for("teachers"))


@app.route("/subjects")
@login_required
@admin_required
def subjects():
    db = get_db()
    rows = db.execute("SELECT s.*, d.name department_name FROM subjects s LEFT JOIN departments d ON d.id=s.department_id ORDER BY s.year, s.code").fetchall()
    db.close()
    grouped = {y: [] for y in YEAR_OPTIONS}
    for row in rows:
        grouped.setdefault(row["year"] or "1st Year", []).append(row)
    return render_template("subjects.html", subjects=rows, grouped=grouped, years=YEAR_OPTIONS)


@app.route("/subjects/add", methods=["GET", "POST"])
@login_required
@admin_required
def add_subject():
    db = get_db()
    departments = db.execute("SELECT * FROM departments WHERE active=1 ORDER BY name").fetchall()
    if request.method == "POST":
        code = request.form.get("code", "").strip().upper(); name = request.form.get("name", "").strip()
        year = request.form.get("year", "").strip()
        # Department is intentionally hidden from the UI. Keep the existing
        # internal department relationship by assigning the default department.
        default_dept_row = db.execute("SELECT id FROM departments ORDER BY id LIMIT 1").fetchone()
        dept_id = default_dept_row["id"] if default_dept_row else None
        if year not in YEAR_OPTIONS: year = ""
        if not code or not name or not year or not dept_id:
            db.close(); flash("Subject code, name and year are required.", "error")
            return render_template("subject_form.html", subject=None, title="Add Subject", departments=departments, years=YEAR_OPTIONS)
        try:
            db.execute("INSERT INTO subjects(code,name,year,department_id) VALUES(?,?,?,?)", (code, name, year, dept_id)); db.commit(); flash("Subject added.", "success")
        except Exception as exc:
            if not isinstance(exc, sqlite3.IntegrityError) and not (POSTGRES_AVAILABLE and isinstance(exc, psycopg2.IntegrityError)): raise
            flash("Subject code already exists.", "error")
        db.close(); return redirect(url_for("subjects"))
    db.close(); return render_template("subject_form.html", subject=None, title="Add Subject", departments=departments, years=YEAR_OPTIONS)


@app.route("/subjects/edit/<int:subject_id>", methods=["GET", "POST"])
@login_required
@admin_required
def edit_subject(subject_id):
    db = get_db(); subject = db.execute("SELECT * FROM subjects WHERE id=?", (subject_id,)).fetchone()
    if not subject:
        db.close(); flash("Subject not found.", "error"); return redirect(url_for("subjects"))
    if request.method == "POST":
        code = request.form.get("code", "").strip().upper()
        name = request.form.get("name", "").strip()
        year = request.form.get("year", "").strip()
        # Preserve the existing internal department relationship; users no longer select it.
        dept_id = subject["department_id"]
        if not code or not name or year not in YEAR_OPTIONS:
            flash("Subject code, name and year are required.", "error")
        else:
            try:
                db.execute("UPDATE subjects SET code=?,name=?,year=?,department_id=? WHERE id=?", (code, name, year, dept_id, subject_id)); db.commit(); flash("Subject updated.", "success"); db.close(); return redirect(url_for("subjects"))
            except Exception as exc:
                if not isinstance(exc, sqlite3.IntegrityError) and not (POSTGRES_AVAILABLE and isinstance(exc, psycopg2.IntegrityError)): raise
                flash("Subject code already exists.", "error")
    departments=db.execute("SELECT * FROM departments WHERE active=1 ORDER BY name").fetchall(); db.close(); return render_template("subject_form.html", subject=subject, title="Edit Subject", departments=departments, years=YEAR_OPTIONS)


@app.route("/subjects/delete/<int:subject_id>", methods=["POST"])
@login_required
@admin_required
def delete_subject(subject_id):
    db = get_db(); db.execute("DELETE FROM subjects WHERE id=?", (subject_id,)); db.commit(); db.close()
    flash("Subject deleted.", "success"); return redirect(url_for("subjects"))


def _prn_sort_key(row):
    value = str(row["prn"] or "").strip()
    if value.isdigit():
        return (0, int(value), "")
    return (1, 0, value.lower())


@app.route("/attendance", methods=["GET", "POST"])
@login_required
@staff_required
def take_attendance():
    db=get_db(); subjects=teacher_or_admin_subjects(db); years=YEAR_OPTIONS[:]
    selected_year=request.args.get('year') or request.form.get('year') or '1st Year'
    session_type=(request.args.get('session_type') or request.form.get('session_type') or 'Lecture').strip()
    if session_type not in ('Lecture','Practical'): session_type='Lecture'
    selected_batch=(request.args.get('batch') or request.form.get('batch') or '').strip()
    selected_subject=int(request.args.get('subject_id') or request.form.get('subject_id') or 0)
    selected_date=request.args.get('attendance_date') or request.form.get('attendance_date') or date.today().isoformat()
    try: lecture_count=max(1,min(3,int(request.args.get('lecture_count') or request.form.get('lecture_count') or 1)))
    except ValueError: lecture_count=1
    if session_type=='Practical' and selected_batch not in ('Batch A','Batch B','Batch C','Batch D'): selected_batch=''
    subjects=[s for s in subjects if s['year']==selected_year]
    if selected_subject and not any(s['id']==selected_subject for s in subjects): selected_subject=0
    if selected_subject and not can_use_subject(db,selected_subject): selected_subject=0
    if session_type=='Practical':
        students=db.execute('SELECT * FROM students WHERE year=? AND batch=? ORDER BY prn,name',(selected_year,selected_batch)).fetchall() if selected_subject and selected_batch else []
    else:
        students=db.execute('SELECT * FROM students WHERE year=? ORDER BY prn,name',(selected_year,)).fetchall() if selected_subject else []
    students=sorted(students,key=_prn_sort_key)
    day_name=date.fromisoformat(selected_date).strftime('%A') if selected_date else date.today().strftime('%A')
    candidates=timetable_candidates(db,selected_year,session_type,selected_batch,selected_subject,day_name) if selected_subject else []
    teachers=all_teacher_rows(db); lecture_data={}
    if request.method=='POST' and selected_subject and students:
        for n in range(1,lecture_count+1):
            tid=request.form.get(f'timetable_id_{n}',type=int); tr=db.execute('SELECT * FROM timetable WHERE id=?',(tid,)).fetchone() if tid else None
            start_time=request.form.get(f'start_time_{n}','').strip(); end_time=request.form.get(f'end_time_{n}','').strip()
            teacher_status=request.form.get(f'teacher_status_{n}','Scheduled teacher present').strip(); conducted=request.form.get(f'conducted_by_{n}',type=int)
            scheduled=tr['teacher_id'] if tr else None
            if tr and (tr['year']!=selected_year or tr['subject_id']!=selected_subject or tr['day_of_week']!=day_name): scheduled=None
            if teacher_status not in ('Scheduled teacher present','Scheduled teacher absent - substitute teacher','Another teacher conducted the session'): teacher_status='Scheduled teacher present'
            if teacher_status=='Scheduled teacher present': conducted=scheduled or session['user_id']
            elif not conducted: conducted=session['user_id']
            if not start_time or not end_time:
                flash(f'Please select start and end time for Lecture {n}.','error'); continue
            lecture_data[n]=(start_time,end_time,scheduled,conducted,teacher_status)
        for st in students:
            for n in range(1,lecture_count+1):
                att_status=request.form.get(f'status_{st["id"]}_{n}')
                if att_status not in ('Present','Absent') or n not in lecture_data: continue
                start_time,end_time,scheduled,conducted,teacher_status=lecture_data[n]
                db.execute("""INSERT INTO attendance(student_id,subject_id,attendance_date,lecture_no,lecture_time,start_time,end_time,session_type,batch,status,marked_by,academic_year,scheduled_teacher_id,conducted_by_teacher_id,teacher_status)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(student_id,subject_id,attendance_date,lecture_no) DO UPDATE SET lecture_time=excluded.lecture_time,start_time=excluded.start_time,end_time=excluded.end_time,session_type=excluded.session_type,batch=excluded.batch,status=excluded.status,marked_by=excluded.marked_by,academic_year=excluded.academic_year,scheduled_teacher_id=excluded.scheduled_teacher_id,conducted_by_teacher_id=excluded.conducted_by_teacher_id,teacher_status=excluded.teacher_status""",
                    (st['id'],selected_subject,selected_date,n,f'{start_time}-{end_time}',start_time,end_time,session_type,selected_batch if session_type=='Practical' else '',att_status,session['user_id'],current_academic_year_value(db),scheduled,conducted,teacher_status))
        db.commit(); flash(f'{session_type} attendance saved with teacher status.','success')
    existing={}; lecture_times={}
    if selected_subject:
        rows=db.execute("SELECT student_id,lecture_no,lecture_time,status FROM attendance WHERE subject_id=? AND attendance_date=? AND COALESCE(session_type,'Lecture')=? AND (COALESCE(batch,'')=? OR ?='Lecture')",(selected_subject,selected_date,session_type,selected_batch,session_type)).fetchall()
        for r in rows: existing[(r['student_id'],r['lecture_no'])]=r['status']; lecture_times.setdefault(r['lecture_no'],r['lecture_time'] or '')
        if rows and request.method=='GET': lecture_count=max(lecture_count,max(r['lecture_no'] for r in rows))
    batch_counts={b:db.execute('SELECT COUNT(*) FROM students WHERE year=? AND batch=?',(selected_year,b)).fetchone()[0] for b in ('Batch A','Batch B','Batch C','Batch D')}
    db.close()
    return render_template('attendance.html',subjects=subjects,students=students,years=years,selected_year=selected_year,selected_subject=selected_subject,selected_date=selected_date,existing=existing,lecture_count=lecture_count,lecture_times=lecture_times,session_type=session_type,selected_batch=selected_batch,batch_options=['Batch A','Batch B','Batch C','Batch D'],batch_counts=batch_counts,candidates=candidates,teachers=teachers,day_name=day_name)


@app.route("/attendance/history")
@login_required
@staff_required
def attendance_history():
    db = get_db()
    where = ["1=1"]; params = []
    if session.get("role") == "teacher":
        where.append("a.subject_id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?)")
        params.append(session["user_id"])
    elif session.get("role") == "hod":
        where.append("sub.department_id=(SELECT department_id FROM users WHERE id=? )")
        params.append(session["user_id"])
    year = request.args.get("year", "").strip()
    session_type = request.args.get("session_type", "").strip()
    batch = request.args.get("batch", "").strip()
    subject_id = request.args.get("subject_id", type=int)
    from_date = request.args.get("from_date", "").strip(); to_date = request.args.get("to_date", "").strip()
    if year:
        where.append("s.year=?"); params.append(year)
    if session_type in ("Lecture","Practical"):
        where.append("COALESCE(a.session_type,'Lecture')=?"); params.append(session_type)
    if batch and session_type == "Practical":
        where.append("COALESCE(a.batch,'')=?"); params.append(batch)
    if subject_id:
        where.append("a.subject_id=?"); params.append(subject_id)
    if from_date:
        where.append("a.attendance_date>=?"); params.append(from_date)
    if to_date:
        where.append("a.attendance_date<=?"); params.append(to_date)
    condition=" AND ".join(where)
    rows=db.execute(f"""SELECT a.attendance_date,a.lecture_no,a.lecture_time,COALESCE(a.session_type,'Lecture') session_type,COALESCE(a.batch,'') batch,sub.id subject_id,sub.code,sub.name subject_name,s.year,
        COUNT(a.id) total, SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END) present,
        SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END) absent
        FROM attendance a JOIN subjects sub ON sub.id=a.subject_id JOIN students s ON s.id=a.student_id
        WHERE {condition}
        GROUP BY a.attendance_date,a.lecture_no,a.lecture_time,a.session_type,a.batch,sub.id,sub.code,sub.name,s.year
        ORDER BY a.attendance_date DESC,a.lecture_no DESC,sub.code,s.year""", params).fetchall()
    subjects=teacher_or_admin_subjects(db); years=YEAR_OPTIONS
    db.close()
    return render_template("attendance_history.html", rows=rows, subjects=subjects, years=years, year=year, subject_id=subject_id, from_date=from_date, to_date=to_date, session_type=session_type, batch=batch, batch_options=["Batch A","Batch B","Batch C","Batch D"])


@app.route("/attendance/history/delete", methods=["POST"])
@login_required
@admin_required
def delete_attendance_session():
    attendance_date=request.form.get("attendance_date", "").strip(); subject_id=request.form.get("subject_id", type=int); lecture_no=request.form.get("lecture_no", type=int)
    if not attendance_date or not subject_id or not lecture_no:
        flash("Invalid attendance session.", "error"); return redirect(url_for("attendance_history"))
    db=get_db()
    subject=db.execute("SELECT code,name FROM subjects WHERE id=?",(subject_id,)).fetchone()
    cur=db.execute("DELETE FROM attendance WHERE subject_id=? AND attendance_date=? AND lecture_no=?",(subject_id,attendance_date,lecture_no))
    deleted=cur.rowcount
    db.commit()
    log_activity("Attendance session deleted", f"{attendance_date} | {subject['code'] if subject else subject_id} | Lecture {lecture_no} | {deleted} records")
    db.close()
    flash("Attendance session deleted successfully.", "success")
    return redirect(url_for("attendance_history"))


@app.route("/students/promotion", methods=["GET","POST"])
@login_required
@admin_required
def student_promotion():
    db=get_db()
    current_row=db.execute("SELECT value FROM app_settings WHERE key='academic_year'").fetchone()
    current_ay=current_row["value"] if current_row and current_row["value"] else "2026-27"
    if request.method=="POST":
        from_year=request.form.get("from_year", "").strip(); to_year=request.form.get("to_year", "").strip(); new_ay=request.form.get("new_academic_year", "").strip()
        if from_year not in YEAR_OPTIONS or to_year not in YEAR_OPTIONS or YEAR_OPTIONS.index(to_year) != YEAR_OPTIONS.index(from_year)+1 or not new_ay:
            flash("Select consecutive years and enter the new academic year.", "error")
        else:
            students=db.execute("SELECT id FROM students WHERE year=?",(from_year,)).fetchall(); changed=0
            for st in students:
                db.execute("INSERT INTO student_academic_history(student_id,academic_year,year,promoted_by) VALUES(?,?,?,?) ON CONFLICT(student_id,academic_year) DO NOTHING", (st["id"],new_ay,to_year,session["user_id"]))
                db.execute("UPDATE students SET year=?,academic_year=? WHERE id=?",(to_year,new_ay,st["id"])); changed+=1
            db.execute("INSERT INTO app_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",("academic_year",new_ay))
            db.commit(); log_activity("Students promoted",f"{from_year} to {to_year} | {new_ay} | {changed} students")
            flash(f"{changed} student(s) promoted to {to_year} for {new_ay}.","success")
            current_ay=new_ay
    counts={y:db.execute("SELECT COUNT(*) FROM students WHERE year=?",(y,)).fetchone()[0] for y in YEAR_OPTIONS}
    history=db.execute("""SELECT h.academic_year,h.year,h.promoted_at,s.prn,s.name FROM student_academic_history h JOIN students s ON s.id=h.student_id ORDER BY h.promoted_at DESC LIMIT 300""").fetchall()
    db.close(); return render_template("promotion.html",years=YEAR_OPTIONS,current_ay=current_ay,counts=counts,history=history)


@app.route("/reports")
@login_required
@staff_required
def reports():
    db = get_db(); subjects = teacher_or_admin_subjects(db); students = db.execute("SELECT * FROM students").fetchall()
    subject_id = request.args.get("subject_id", type=int); student_id = request.args.get("student_id", type=int)
    year = request.args.get("year", "")
    session_type = request.args.get("session_type", "")
    batch = request.args.get("batch", "")
    if session_type not in ("", "Lecture", "Practical"):
        session_type = ""
    if batch not in ("", "Batch A", "Batch B", "Batch C", "Batch D"):
        batch = ""
    if session_type != "Practical":
        batch = ""
    from_date = request.args.get("from_date", ""); to_date = request.args.get("to_date", "")
    if subject_id and not can_use_subject(db, subject_id):
        subject_id = None; flash("You are not assigned to that subject.", "error")
    on_where = ["1=1"]; on_params = []
    student_where = ["1=1"]; student_params = []
    if session.get("role") == "teacher":
        on_where.append("a.subject_id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?)"); on_params.append(session["user_id"])
    elif session.get("role") == "hod":
        on_where.append("a.subject_id IN (SELECT id FROM subjects WHERE department_id=(SELECT department_id FROM users WHERE id=?))"); on_params.append(session["user_id"])
    if subject_id: on_where.append("a.subject_id=?"); on_params.append(subject_id)
    if session_type: on_where.append("COALESCE(a.session_type,'Lecture')=?"); on_params.append(session_type)
    if session_type == "Practical" and batch: on_where.append("COALESCE(a.batch,'')=?"); on_params.append(batch)
    if from_date: on_where.append("a.attendance_date>=?"); on_params.append(from_date)
    if to_date: on_where.append("a.attendance_date<=?"); on_params.append(to_date)
    if student_id: student_where.append("s.id=?"); student_params.append(student_id)
    if year: student_where.append("s.year=?"); student_params.append(year)
    if session_type == "Practical" and batch: student_where.append("COALESCE(s.batch,'')=?"); student_params.append(batch)
    on_condition = " AND ".join(on_where)
    student_condition = " AND ".join(student_where)
    summary = db.execute(f"""SELECT s.id,s.prn,s.name,s.year,s.course,COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM students s LEFT JOIN attendance a ON a.student_id=s.id AND {on_condition}
        WHERE {student_condition}
        GROUP BY s.id""", on_params + student_params).fetchall()
    summary = sorted(summary, key=_prn_sort_key)
    students = sorted(students, key=_prn_sort_key)
    detail_where = ["1=1"]; detail_params = []
    if session.get("role") == "teacher":
        detail_where.append("a.subject_id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?)"); detail_params.append(session["user_id"])
    elif session.get("role") == "hod":
        detail_where.append("a.subject_id IN (SELECT id FROM subjects WHERE department_id=(SELECT department_id FROM users WHERE id=?))"); detail_params.append(session["user_id"])
    if subject_id: detail_where.append("a.subject_id=?"); detail_params.append(subject_id)
    if session_type: detail_where.append("COALESCE(a.session_type,'Lecture')=?"); detail_params.append(session_type)
    if session_type == "Practical" and batch: detail_where.append("COALESCE(a.batch,'')=?"); detail_params.append(batch)
    if student_id: detail_where.append("a.student_id=?"); detail_params.append(student_id)
    if year: detail_where.append("s.year=?"); detail_params.append(year)
    if from_date: detail_where.append("a.attendance_date>=?"); detail_params.append(from_date)
    if to_date: detail_where.append("a.attendance_date<=?"); detail_params.append(to_date)
    detail_condition = " AND ".join(detail_where)
    detail = db.execute(f"""SELECT a.attendance_date,a.lecture_no,a.lecture_time,s.prn,s.name,sub.code,sub.name subject,a.status
        FROM attendance a JOIN students s ON s.id=a.student_id JOIN subjects sub ON sub.id=a.subject_id
        WHERE {detail_condition} ORDER BY a.attendance_date DESC,a.lecture_no""", detail_params).fetchall()
    detail = sorted(detail, key=_prn_sort_key)
    detail = sorted(detail, key=lambda r: (str(r["attendance_date"] or ""), int(r["lecture_no"] or 0)), reverse=True)
    years = [r["year"] for r in db.execute("SELECT DISTINCT year FROM students WHERE year IS NOT NULL AND year<>'' ORDER BY year").fetchall()]
    subject_scope_where = "1=1"
    subject_scope_params = []
    if session.get("role") == "teacher":
        subject_scope_where = "sub.id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?)"
        subject_scope_params = [session["user_id"]]
    elif session.get("role") == "hod":
        subject_scope_where = "sub.department_id=(SELECT department_id FROM users WHERE id=?)"
        subject_scope_params = [session["user_id"]]
    subject_summary = db.execute(f"""SELECT sub.code,sub.name,COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM subjects sub LEFT JOIN attendance a ON a.subject_id=sub.id AND {on_condition}
        WHERE {subject_scope_where}
        GROUP BY sub.id ORDER BY sub.code""", on_params + subject_scope_params).fetchall()
    for y in YEAR_OPTIONS:
        if y not in years: years.append(y)
    db.close()
    return render_template("reports.html", subjects=subjects, students=students, years=years, summary=summary, detail=detail, subject_summary=subject_summary,
                           subject_id=subject_id, student_id=student_id, year=year, session_type=session_type, batch=batch, batch_options=["Batch A","Batch B","Batch C","Batch D"], from_date=from_date, to_date=to_date)


@app.route("/reports/export.csv")
@login_required
@staff_required
def export_csv():
    db = get_db(); subject_id = request.args.get("subject_id", type=int); student_id = request.args.get("student_id", type=int)
    year = request.args.get("year", "")
    session_type = request.args.get("session_type", "")
    batch = request.args.get("batch", "")
    if session_type not in ("", "Lecture", "Practical"):
        session_type = ""
    if batch not in ("", "Batch A", "Batch B", "Batch C", "Batch D"):
        batch = ""
    if session_type != "Practical":
        batch = ""
    from_date = request.args.get("from_date", ""); to_date = request.args.get("to_date", "")
    on_where = ["1=1"]; on_params = []
    student_where = ["1=1"]; student_params = []
    if session.get("role") != "admin":
        on_where.append("a.subject_id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?)"); on_params.append(session["user_id"])
    if subject_id and can_use_subject(db, subject_id): on_where.append("a.subject_id=?"); on_params.append(subject_id)
    if session_type: on_where.append("COALESCE(a.session_type,'Lecture')=?"); on_params.append(session_type)
    if session_type == "Practical" and batch: on_where.append("COALESCE(a.batch,'')=?"); on_params.append(batch)
    if from_date: on_where.append("a.attendance_date>=?"); on_params.append(from_date)
    if to_date: on_where.append("a.attendance_date<=?"); on_params.append(to_date)
    if student_id: student_where.append("s.id=?"); student_params.append(student_id)
    if year: student_where.append("s.year=?"); student_params.append(year)
    if session_type == "Practical" and batch: student_where.append("COALESCE(s.batch,'')=?"); student_params.append(batch)
    on_condition = " AND ".join(on_where)
    student_condition = " AND ".join(student_where)
    rows = db.execute(f"""SELECT s.prn,s.name,s.course,s.year,
        COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM students s LEFT JOIN attendance a ON a.student_id=s.id AND {on_condition}
        WHERE {student_condition}
        GROUP BY s.id""", on_params + student_params).fetchall()
    subject_summary = db.execute(f"""SELECT sub.code,sub.name,COUNT(a.id) total, COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present, COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent, CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct FROM subjects sub LEFT JOIN attendance a ON a.subject_id=sub.id AND {on_condition} GROUP BY sub.id ORDER BY sub.code""", on_params).fetchall()
    db.close(); out = io.StringIO(); writer = csv.writer(out)
    writer.writerow(["SUBJECT-WISE SUMMARY"])
    writer.writerow(["Subject","Total Attendance","Present","Absent","Attendance Percentage"])
    for r in subject_summary: writer.writerow([f"{r['code']} - {r['name']}",r['total'],r['present'],r['absent'],f"{r['pct']}%"] )
    writer.writerow([])
    writer.writerow(["STUDENT SUMMARY"])
    writer.writerow(["PRN","Name","Year","Course","Total Attendance","Present Days","Absent Days","Attendance Percentage"])
    for r in rows:
        writer.writerow([r["prn"], r["name"], r["year"], r["course"], r["total"], r["present"], r["absent"], f'{r["pct"]}%'])
    return Response(out.getvalue(), mimetype="text/csv", headers={"Content-Disposition":"attachment; filename=attendance_summary.csv"})


@app.route("/reports/export.pdf")
@login_required
@staff_required
def export_pdf():
    db = get_db()
    subject_id = request.args.get("subject_id", type=int)
    student_id = request.args.get("student_id", type=int)
    year = request.args.get("year", "")
    session_type = request.args.get("session_type", "")
    batch = request.args.get("batch", "")
    if session_type not in ("", "Lecture", "Practical"):
        session_type = ""
    if batch not in ("", "Batch A", "Batch B", "Batch C", "Batch D"):
        batch = ""
    if session_type != "Practical":
        batch = ""
    from_date = request.args.get("from_date", "")
    to_date = request.args.get("to_date", "")

    on_where = ["1=1"]; on_params = []
    student_where = ["1=1"]; student_params = []
    if session.get("role") != "admin":
        on_where.append("a.subject_id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?)")
        on_params.append(session["user_id"])
    if subject_id and can_use_subject(db, subject_id):
        on_where.append("a.subject_id=?"); on_params.append(subject_id)
    elif subject_id:
        subject_id = None
    if session_type:
        on_where.append("COALESCE(a.session_type,'Lecture')=?"); on_params.append(session_type)
    if session_type == "Practical" and batch:
        on_where.append("COALESCE(a.batch,'')=?"); on_params.append(batch)
    if from_date:
        on_where.append("a.attendance_date>=?"); on_params.append(from_date)
    if to_date:
        on_where.append("a.attendance_date<=?"); on_params.append(to_date)
    if student_id:
        student_where.append("s.id=?"); student_params.append(student_id)
    if year:
        student_where.append("s.year=?"); student_params.append(year)
    if session_type == "Practical" and batch:
        student_where.append("COALESCE(s.batch,'')=?"); student_params.append(batch)

    on_condition = " AND ".join(on_where)
    student_condition = " AND ".join(student_where)
    rows = db.execute(f"""SELECT s.prn,s.name,s.course,s.year,
        COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM students s LEFT JOIN attendance a ON a.student_id=s.id AND {on_condition}
        WHERE {student_condition}
        GROUP BY s.id""", on_params + student_params).fetchall()
    rows = sorted(rows, key=_prn_sort_key)

    subject_name = "All Subjects"
    if subject_id:
        subject_row = db.execute("SELECT code,name FROM subjects WHERE id=?", (subject_id,)).fetchone()
        if subject_row:
            subject_name = f"{subject_row['code']} - {subject_row['name']}"

    teacher_rows = db.execute(f"SELECT DISTINCT u.full_name,u.username FROM attendance a LEFT JOIN users u ON u.id=a.scheduled_teacher_id WHERE {on_condition} AND a.scheduled_teacher_id IS NOT NULL ORDER BY u.full_name,u.username", on_params).fetchall()
    conducted_rows = db.execute(f"SELECT DISTINCT u.full_name,u.username FROM attendance a LEFT JOIN users u ON u.id=a.conducted_by_teacher_id WHERE {on_condition} AND a.conducted_by_teacher_id IS NOT NULL ORDER BY u.full_name,u.username", on_params).fetchall()
    scheduled_teacher_text = ", ".join((r["full_name"] or r["username"]) for r in teacher_rows) or "—"
    conducted_teacher_text = ", ".join((r["full_name"] or r["username"]) for r in conducted_rows) or "—"
    db.close()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        rightMargin=10*mm, leftMargin=10*mm, topMargin=8*mm, bottomMargin=12*mm
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "CollegeTitle", parent=styles["Title"], alignment=TA_CENTER,
        fontSize=21, leading=24, spaceAfter=1, textColor=colors.HexColor("#12355b")
    )
    affiliation_style = ParagraphStyle(
        "Affiliation", parent=styles["Normal"], alignment=TA_CENTER,
        fontSize=12.5, leading=16, spaceAfter=1, textColor=colors.HexColor("#15539a")
    )
    course_style = ParagraphStyle(
        "Course", parent=styles["Normal"], alignment=TA_CENTER,
        fontSize=11.5, leading=14, spaceAfter=2, textColor=colors.HexColor("#12355b")
    )
    report_style = ParagraphStyle(
        "ReportSub", parent=styles["Normal"], alignment=TA_CENTER,
        fontSize=18, leading=21, spaceAfter=2, textColor=colors.white
    )
    meta_style = ParagraphStyle(
        "ReportMeta", parent=styles["Normal"], alignment=TA_LEFT,
        fontSize=10.5, leading=14, textColor=colors.HexColor("#12355b")
    )
    small_style = ParagraphStyle(
        "Small", parent=styles["Normal"], fontSize=8.5, leading=11,
        textColor=colors.HexColor("#334155")
    )

    story = []
    logo_path = BASE_DIR / "static" / LOGO_FILE
    logo = Image(str(logo_path), width=28*mm, height=28*mm) if logo_path.exists() else ""

    class_name = year if year else "All Classes"
    course_name = "B.Pharm"
    teacher_rows = db.execute(f"SELECT DISTINCT u.full_name,u.username FROM attendance a LEFT JOIN users u ON u.id=a.scheduled_teacher_id WHERE {on_condition} AND a.scheduled_teacher_id IS NOT NULL ORDER BY u.full_name,u.username", on_params).fetchall()
    conducted_rows = db.execute(f"SELECT DISTINCT u.full_name,u.username FROM attendance a LEFT JOIN users u ON u.id=a.conducted_by_teacher_id WHERE {on_condition} AND a.conducted_by_teacher_id IS NOT NULL ORDER BY u.full_name,u.username", on_params).fetchall()
    scheduled_teacher_text = ", ".join((r["full_name"] or r["username"]) for r in teacher_rows) or "—"
    conducted_teacher_text = ", ".join((r["full_name"] or r["username"]) for r in conducted_rows) or "—"

    # FIRST PAGE ONLY: full institutional header.
    header_text = [
        Paragraph(xml_escape(str(COLLEGE_NAME)), title_style),
        Paragraph(
            "Affiliated by Dr. Babasaheb Ambedkar Technological University, Lonere, Raigad",
            affiliation_style
        ),
        Paragraph("B.Pharm", course_style),
    ]
    header = Table([[logo, header_text]], colWidths=[32*mm, 230*mm])
    header.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ALIGN", (1,0), (1,0), "CENTER"),
        ("LEFTPADDING", (0,0), (-1,-1), 0),
        ("RIGHTPADDING", (0,0), (-1,-1), 0),
        ("TOPPADDING", (0,0), (-1,-1), 0),
        ("BOTTOMPADDING", (0,0), (-1,-1), 0),
    ]))
    story.append(header)
    story.append(Spacer(1, 3*mm))

    # Attendance report title ribbon.
    report_banner = Table(
        [[Paragraph("▣ &nbsp; ATTENDANCE REPORT", report_style)]],
        colWidths=[262*mm]
    )
    report_banner.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#07579a")),
        ("BOX", (0,0), (-1,-1), 0.5, colors.HexColor("#07579a")),
        ("LEFTPADDING", (0,0), (-1,-1), 6),
        ("RIGHTPADDING", (0,0), (-1,-1), 6),
        ("TOPPADDING", (0,0), (-1,-1), 7),
        ("BOTTOMPADDING", (0,0), (-1,-1), 7),
    ]))
    story.append(report_banner)
    story.append(Spacer(1, 4*mm))

    # Information strip.
    period_text = f"{from_date or '—'} – {to_date or '—'}"
    info_data = [[
        Paragraph(f"<b>Subject</b>&nbsp;&nbsp;: &nbsp;{xml_escape(str(subject_name))}", meta_style),
        Paragraph(f"<b>Class</b>&nbsp;&nbsp;: &nbsp;{xml_escape(str(class_name))}", meta_style),
        Paragraph(f"<b>Date Range</b>&nbsp;&nbsp;: &nbsp;{period_text}", meta_style),
    ],[
        Paragraph(f"<b>Scheduled Teacher</b>&nbsp;&nbsp;: &nbsp;{xml_escape(scheduled_teacher_text)}", meta_style),
        Paragraph(f"<b>Conducted By</b>&nbsp;&nbsp;: &nbsp;{xml_escape(conducted_teacher_text)}", meta_style),
        Paragraph(f"<b>Session</b>&nbsp;&nbsp;: &nbsp;{xml_escape(session_type or 'All')}{(' / '+xml_escape(batch)) if batch else ''}", meta_style),
    ]]
    info = Table(info_data, colWidths=[88*mm, 70*mm, 104*mm])
    info.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#eaf4ff")),
        ("BOX", (0,0), (-1,-1), 0.5, colors.HexColor("#c6def4")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 7),
        ("RIGHTPADDING", (0,0), (-1,-1), 7),
        ("TOPPADDING", (0,0), (-1,-1), 8),
        ("BOTTOMPADDING", (0,0), (-1,-1), 8),
    ]))
    story.append(info)
    story.append(Spacer(1, 5*mm))

    # Only the table header repeats on subsequent pages.
    data = [["PRN", "Student Name", "Total Lectures", "Present", "Absent", "Attendance %"]]
    for r in rows:
        data.append([
            str(r["prn"] or "—"), str(r["name"] or ""), r["total"], r["present"],
            r["absent"], f'{r["pct"]}%'
        ])

    table = Table(
        data, repeatRows=1,
        colWidths=[38*mm, 67*mm, 36*mm, 32*mm, 32*mm, 40*mm],
        hAlign="CENTER"
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#07579a")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTNAME", (0,1), (-1,-1), "Helvetica"),
        ("FONTSIZE", (0,0), (-1,0), 9.5),
        ("FONTSIZE", (0,1), (-1,-1), 8.8),
        ("LEADING", (0,0), (-1,-1), 11),
        ("ALIGN", (0,0), (-1,-1), "CENTER"),
        ("ALIGN", (1,1), (1,-1), "LEFT"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("GRID", (0,0), (-1,-1), 0.45, colors.HexColor("#9fc5e8")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1),
         [colors.white, colors.HexColor("#eef7ff")]),
        ("TOPPADDING", (0,0), (-1,-1), 5.5),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5.5),
    ]))
    story.append(table)

    # Signature is after the table, so it appears only at the end of the report.
    story.append(Spacer(1, 7*mm))
    generated = datetime.now(ZoneInfo("Asia/Kolkata")).strftime("%d %b %Y | %I:%M %p")
    sign_data = [[
        Paragraph(f"<b>Generated On</b>&nbsp;&nbsp;: &nbsp;{generated}", small_style),
        Paragraph(
            "______________________________<br/><b>Teacher Signature</b><br/>(Subject Incharge)",
            small_style
        )
    ]]
    sign = Table(sign_data, colWidths=[130*mm, 132*mm])
    sign.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "BOTTOM"),
        ("ALIGN", (1,0), (1,0), "CENTER"),
        ("LEFTPADDING", (0,0), (-1,-1), 0),
        ("RIGHTPADDING", (0,0), (-1,-1), 0),
        ("TOPPADDING", (0,0), (-1,-1), 0),
        ("BOTTOMPADDING", (0,0), (-1,-1), 0),
    ]))
    story.append(sign)

    def draw_page_footer(canvas, doc_obj):
        canvas.saveState()
        page_w, page_h = landscape(A4)

        # Dark blue footer band.
        footer_h = 9*mm
        canvas.setFillColor(colors.HexColor("#084b83"))
        canvas.rect(0, 0, page_w, footer_h, fill=1, stroke=0)

        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica", 9)
        canvas.drawCentredString(
            page_w / 2, 3.3*mm, "Better Education  •  Brighter Future"
        )
        canvas.drawRightString(
            page_w - 10*mm, 3.3*mm, f"Page {doc_obj.page}"
        )

        # For pages after page 1, keep only a compact report context line.
        if doc_obj.page > 1:
            canvas.setFillColor(colors.HexColor("#12355b"))
            canvas.setFont("Helvetica-Bold", 8.5)
            canvas.drawString(
                10*mm, page_h - 5*mm,
                f"ATTENDANCE REPORT  |  {subject_name}  |  {class_name}"
            )
            canvas.setStrokeColor(colors.HexColor("#c6def4"))
            canvas.line(10*mm, page_h - 6.5*mm, page_w - 10*mm, page_h - 6.5*mm)

        canvas.restoreState()

    # The complete college header is deliberately NOT repeated.
    doc.build(story, onFirstPage=draw_page_footer, onLaterPages=draw_page_footer)

    pdf = buffer.getvalue()
    buffer.close()
    return Response(
        pdf,
        mimetype="application/pdf",
        headers={"Content-Disposition":"attachment; filename=attendance_report.pdf"}
    )



@app.route("/activity-log")
@login_required
@admin_required
def activity_log():
    db=get_db(); q=request.args.get("q","").strip(); params=[]
    where="1=1"
    if q: where="(l.action LIKE ? OR l.details LIKE ? OR u.username LIKE ?)"; params=[f"%{q}%"]*3
    rows=db.execute(f"SELECT l.created_at,l.action,l.details,COALESCE(u.username,'System') username FROM activity_logs l LEFT JOIN users u ON u.id=l.user_id WHERE {where} ORDER BY l.created_at DESC LIMIT 300",params).fetchall(); db.close()
    return render_template("activity_log.html",rows=rows,q=q)

@app.route("/profile", methods=["GET","POST"])
@login_required
def profile():
    db=get_db(); user=db.execute("SELECT * FROM users WHERE id=?",(session["user_id"],)).fetchone()
    if request.method=="POST":
        full_name=request.form.get("full_name","").strip(); mobile=request.form.get("mobile","").strip(); new_password=request.form.get("new_password","")
        if new_password:
            err=validate_password(new_password)
            if err: flash(err,"error")
            else:
                db.execute("UPDATE users SET full_name=?,mobile=?,password=? WHERE id=?",(full_name,mobile,generate_password_hash(new_password),session["user_id"])); db.commit(); session["full_name"]=full_name; log_activity("Profile updated","Name/mobile/password updated"); flash("Profile updated.","success")
        else:
            db.execute("UPDATE users SET full_name=?,mobile=? WHERE id=?",(full_name,mobile,session["user_id"])); db.commit(); session["full_name"]=full_name; log_activity("Profile updated","Name/mobile updated"); flash("Profile updated.","success")
        user=db.execute("SELECT * FROM users WHERE id=?",(session["user_id"],)).fetchone()
    db.close(); return render_template("profile.html",user=user)

@app.route("/settings", methods=["GET","POST"])
@login_required
@admin_required
def settings():
    db=get_db()
    if request.method=="POST":
        threshold=request.form.get("threshold","75").strip()
        academic_year=request.form.get("academic_year","").strip()
        if threshold.isdigit() and 0 <= int(threshold) <= 100:
            for k,v in (("attendance_threshold",threshold),("academic_year",academic_year)):
                db.execute("INSERT INTO app_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(k,v))
            db.commit(); log_activity("Settings updated",f"Threshold {threshold}, academic year {academic_year}"); flash("Settings saved.","success")
        else: flash("Threshold must be between 0 and 100.","error")
    vals={r["key"]:r["value"] for r in db.execute("SELECT key,value FROM app_settings").fetchall()}; db.close()
    return render_template("settings.html",settings=vals)

@app.route("/reports/export.xlsx")
@login_required
@staff_required
def export_xlsx():
    if not OPENPYXL_AVAILABLE:
        return Response("Excel export is not installed on this server.", status=500, mimetype="text/plain")
    db = get_db(); subject_id=request.args.get("subject_id", type=int); student_id=request.args.get("student_id", type=int)
    year=request.args.get("year", "")
    session_type=request.args.get("session_type", "")
    batch=request.args.get("batch", "")
    if session_type not in ("", "Lecture", "Practical"):
        session_type = ""
    if batch not in ("", "Batch A", "Batch B", "Batch C", "Batch D"):
        batch = ""
    if session_type != "Practical":
        batch = ""; from_date=request.args.get("from_date", ""); to_date=request.args.get("to_date", "")
    on_where=["1=1"]; on_params=[]; student_where=["1=1"]; student_params=[]
    if session.get("role") != "admin": on_where.append("a.subject_id IN (SELECT subject_id FROM subject_teachers WHERE teacher_id=?)"); on_params.append(session["user_id"])
    if subject_id and can_use_subject(db, subject_id): on_where.append("a.subject_id=?"); on_params.append(subject_id)
    elif subject_id: subject_id=None
    if session_type: on_where.append("COALESCE(a.session_type,'Lecture')=?"); on_params.append(session_type)
    if session_type == "Practical" and batch: on_where.append("COALESCE(a.batch,'')=?"); on_params.append(batch)
    if from_date: on_where.append("a.attendance_date>=?"); on_params.append(from_date)
    if to_date: on_where.append("a.attendance_date<=?"); on_params.append(to_date)
    if student_id: student_where.append("s.id=?"); student_params.append(student_id)
    if year: student_where.append("s.year=?"); student_params.append(year)
    if session_type == "Practical" and batch: student_where.append("COALESCE(s.batch,'')=?"); student_params.append(batch)
    subject_name = "All Subjects"
    if subject_id:
        sr = db.execute("SELECT code,name FROM subjects WHERE id=?", (subject_id,)).fetchone()
        if sr: subject_name = f"{sr['code']} - {sr['name']}"
    rows=db.execute(f"""SELECT s.prn,s.name,s.year,s.course,COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM students s LEFT JOIN attendance a ON a.student_id=s.id AND {' AND '.join(on_where)}
        WHERE {' AND '.join(student_where)} GROUP BY s.id""", on_params+student_params).fetchall()
    subject_summary = db.execute(f"""SELECT sub.code,sub.name,COUNT(a.id) total, COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present, COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent, CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct FROM subjects sub LEFT JOIN attendance a ON a.subject_id=sub.id AND {' AND '.join(on_where)} GROUP BY sub.id ORDER BY sub.code""", on_params).fetchall()
    db.close()
    wb=Workbook(); ws=wb.active; ws.title="Attendance Summary"
    sws=wb.create_sheet("Subject Summary")
    sws.append(["Subject","Total Attendance","Present","Absent","Attendance Percentage"])
    for r in subject_summary: sws.append([f"{r['code']} - {r['name']}",r['total'],r['present'],r['absent'],f"{r['pct']}%"] )
    for c in sws[1]: c.font=Font(bold=True)
    for col,width in zip("ABCDE",[35,20,15,15,25]): sws.column_dimensions[col].width=width

    ws.row_dimensions[1].height = 72
    ws.merge_cells("B1:H1")
    ws["B1"] = COLLEGE_NAME
    ws["B1"].font = Font(bold=True, size=22)
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("B2:H2")
    ws["B2"] = "Attendance Report"
    ws["B2"].font = Font(bold=True, size=13)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("B3:H3")
    ws["B3"] = f"Subject: {subject_name}"
    ws["B3"].font = Font(size=12)
    ws["B3"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("B4:H4")
    ws["B4"] = f"Class: {year if year else 'All Classes'}"
    ws["B4"].font = Font(size=12)
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center")
    filter_text = []
    if from_date or to_date: filter_text.append(f"Date: {from_date or '—'} to {to_date or '—'}")
    if student_id: filter_text.append("Student filter applied")
    if filter_text:
        ws.merge_cells("B5:H5")
        ws["B5"] = " | ".join(filter_text)
        ws["B5"].font = Font(size=10)
        ws["B5"].alignment = Alignment(horizontal="left", vertical="center")
    header_row = 7
    if XLImage is not None:
        logo_path = BASE_DIR / "static" / LOGO_FILE
        if logo_path.exists():
            xl_logo = XLImage(str(logo_path))
            xl_logo.width = 95
            xl_logo.height = 95
            ws.add_image(xl_logo, "A1")
    headers=["PRN","Name","Year","Course","Total Attendance","Present Days","Absent Days","Attendance Percentage"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col_idx, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, r in enumerate(rows, header_row + 1):
        values=[r["prn"],r["name"],r["year"],r["course"],r["total"],r["present"],r["absent"],f'{r["pct"]}%']
        for c_idx, value in enumerate(values, 1):
            ws.cell(r_idx, c_idx, value)
    for col,width in zip("ABCDEFGH",[15,28,16,16,20,15,15,25]): ws.column_dimensions[col].width=width
    ws.row_dimensions[header_row].height = 30
    for row in ws.iter_rows(min_row=header_row):
        for cell in row: cell.alignment=Alignment(vertical="center")
    ws.freeze_panes = f"A{header_row+1}"
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return Response(out.getvalue(),mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=attendance_summary.xlsx"})


@app.route("/reports/print")
@login_required
@staff_required
def print_report():
    return redirect(url_for("reports", **request.args.to_dict()))


@app.route("/admin/backup")
@login_required
@admin_required
def backup_database():
    db=get_db(); tables=["departments","students","users","subjects","subject_teachers","attendance","timetable","notifications","leave_applications","qr_sessions","app_settings"]
    payload={"college":COLLEGE_NAME,"generated_at":datetime.utcnow().isoformat()+"Z","tables":{}}
    for table in tables:
        rows=db.execute(f"SELECT * FROM {table}").fetchall()
        payload["tables"][table]=[dict(r) for r in rows]
    db.close()
    out=io.BytesIO()
    with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
        z.writestr("attendance_backup.json",json.dumps(payload,indent=2,default=str))
        z.writestr("README.txt",f"Y.N.P. College of Pharmacy Attendance backup\nGenerated: {payload['generated_at']}\n\nThis backup contains application data from all attendance tables. Keep it securely because it includes account and attendance information.\n")
    out.seek(0)
    return Response(out.getvalue(),mimetype="application/zip",headers={"Content-Disposition":"attachment; filename=ynp_attendance_backup.zip"})


@app.route("/student-report/<int:student_id>")
@app.route("/student/report")
@login_required
def student_report(student_id=None):
    # Staff can view a selected student's report; students can view their own report.
    db = get_db()
    role = session.get("role")
    if role == "student":
        student = current_student(db)
        if not student:
            db.close()
            flash("Student account is not linked correctly. Contact the administrator.", "error")
            return redirect(url_for("login"))
        student_id = student["id"]
    elif role in ("admin", "teacher", "hod"):
        if student_id is None:
            db.close()
            return redirect(url_for("students"))
        student = db.execute("SELECT * FROM students WHERE id=?", (student_id,)).fetchone()
        if not student:
            db.close(); flash("Student not found.", "error"); return redirect(url_for("students"))
    else:
        db.close()
        return redirect(url_for("dashboard"))
    rows = db.execute("""SELECT sub.code,sub.name,COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM subjects sub LEFT JOIN attendance a ON a.subject_id=sub.id AND a.student_id=?
        GROUP BY sub.id ORDER BY sub.code""", (student_id,)).fetchall()
    db.close(); return render_template("student_report.html", student=student, rows=rows)


@app.route("/student/report/export.pdf")
@login_required
def student_report_export_pdf():
    """Export the logged-in student's attendance report as a professional PDF."""
    if session.get("role") != "student":
        return Response("Not authorized.", status=403)
    db = get_db()
    student = current_student(db)
    if not student:
        db.close()
        return Response("Student record not found.", status=404)
    rows = db.execute("""SELECT sub.code,sub.name,COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        COALESCE(SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END),0) absent,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM subjects sub LEFT JOIN attendance a ON a.subject_id=sub.id AND a.student_id=?
        WHERE sub.department_id=? AND sub.year=?
        GROUP BY sub.id ORDER BY sub.code""",
        (student["id"], student["department_id"], student["year"])).fetchall()
    total = sum(int(r["total"] or 0) for r in rows)
    present = sum(int(r["present"] or 0) for r in rows)
    absent = sum(int(r["absent"] or 0) for r in rows)
    overall = round((present / total) * 100, 1) if total else 0
    db.close()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, rightMargin=14*mm, leftMargin=14*mm,
        topMargin=12*mm, bottomMargin=12*mm
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle('SRCollegeTitle', parent=styles['Title'], alignment=TA_CENTER, fontSize=18, leading=21, spaceAfter=1)
    sub = ParagraphStyle('SRSub', parent=styles['Normal'], alignment=TA_CENTER, fontSize=10, leading=13, spaceAfter=1)
    meta = ParagraphStyle('SRMeta', parent=styles['Normal'], alignment=TA_CENTER, fontSize=8.5, leading=11)
    label = ParagraphStyle('SRLabel', parent=styles['Normal'], fontSize=8.5, leading=11)
    cell = ParagraphStyle('SRCell', parent=styles['Normal'], fontSize=8.2, leading=10)
    head = ParagraphStyle('SRHead', parent=cell, fontName='Helvetica-Bold', alignment=TA_CENTER)

    story = []
    logo_path = BASE_DIR/'static'/LOGO_FILE
    logo = Image(str(logo_path), width=22*mm, height=22*mm) if logo_path.exists() else ''
    header = [
        Paragraph(COLLEGE_NAME, title),
        Paragraph('Student Attendance Report', sub),
        Paragraph('Affiliated by Dr. Babasaheb Ambedkar Technological University, Lonere, Raigad', meta),
    ]
    hdr = Table([[logo, header]], colWidths=[27*mm, 145*mm])
    hdr.setStyle(TableStyle([
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'), ('LEFTPADDING',(0,0),(-1,-1),0),
        ('RIGHTPADDING',(0,0),(-1,-1),0), ('TOPPADDING',(0,0),(-1,-1),0),
        ('BOTTOMPADDING',(0,0),(-1,-1),0),
    ]))
    story.append(hdr)
    story.append(Spacer(1, 5*mm))

    info = Table([
        [Paragraph('<b>Student Name</b><br/>'+str(student['name'] or '—'), label),
         Paragraph('<b>PRN</b><br/>'+str(student['prn'] or '—'), label),
         Paragraph('<b>Year</b><br/>'+str(student['year'] or '—'), label)],
        [Paragraph('<b>Course</b><br/>'+str(student['course'] or '—'), label),
         Paragraph('<b>Overall Attendance</b><br/>'+str(overall)+'%', label),
         Paragraph('<b>Generated</b><br/>'+datetime.now(ZoneInfo('Asia/Kolkata')).strftime('%d-%m-%Y %I:%M %p IST'), label)]
    ], colWidths=[57*mm,57*mm,57*mm])
    info.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#f4f8fd')),
        ('BOX',(0,0),(-1,-1),0.5,colors.HexColor('#c9d8e8')),
        ('INNERGRID',(0,0),(-1,-1),0.35,colors.HexColor('#d8e3ee')),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),
        ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),
    ]))
    story.append(info)
    story.append(Spacer(1, 5*mm))

    data = [[Paragraph('Subject',head), Paragraph('Total',head), Paragraph('Present',head), Paragraph('Absent',head), Paragraph('Attendance',head)]]
    for r in rows:
        data.append([
            Paragraph(f"{r['code']} — {r['name']}", cell),
            Paragraph(str(r['total']), cell), Paragraph(str(r['present']), cell),
            Paragraph(str(r['absent']), cell), Paragraph(f"{r['pct']}%", cell)
        ])
    if not rows:
        data.append([Paragraph('No attendance records found for this student.', cell), '', '', '', ''])
    table = Table(data, repeatRows=1, colWidths=[80*mm,22*mm,24*mm,22*mm,28*mm])
    table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0b3974')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('GRID',(0,0),(-1,-1),0.4,colors.HexColor('#b8c7d9')),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('ALIGN',(1,1),(-1,-1),'CENTER'),
        ('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),
        ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),
    ]))
    story.append(table)
    story.append(Spacer(1, 5*mm))
    story.append(Paragraph(
        f'<b>Total Attendance:</b> {total} &nbsp;&nbsp; <b>Present:</b> {present} &nbsp;&nbsp; <b>Absent:</b> {absent} &nbsp;&nbsp; <b>Overall:</b> {overall}%',
        meta
    ))
    story.append(Spacer(1, 3*mm))
    story.append(Paragraph('Y.N.P. College of Pharmacy — Student Attendance Report', meta))
    doc.build(story)
    buffer.seek(0)
    filename = f"student_report_{str(student['prn'] or student['id']).replace(' ','_')}.pdf"
    return Response(buffer.getvalue(), mimetype='application/pdf', headers={'Content-Disposition': f'attachment; filename={filename}'})



def notify_user(db, user_id, title, message, link=""):
    if user_id:
        db.execute("INSERT INTO notifications(user_id,title,message,link,created_at) VALUES(?,?,?,?,?)",
                   (user_id, title, message, link, datetime.utcnow().isoformat()))


def notify_students_for_attendance(db, subject_id, year, attendance_date, lecture_no):
    rows = db.execute("""SELECT u.id,u.student_id,s.name,sub.name subject FROM users u
        JOIN students s ON s.id=u.student_id JOIN subjects sub ON sub.id=?
        WHERE u.role='student' AND u.approved=1 AND s.year=?""", (subject_id,year)).fetchall()
    for r in rows:
        notify_user(db, r["id"], "Attendance Updated", f"Attendance was recorded for {r['subject']} (Lecture {lecture_no}) on {attendance_date}.", "/student/dashboard")


def create_low_attendance_notifications(db):
    threshold_row = db.execute("SELECT value FROM app_settings WHERE key='attendance_threshold'").fetchone()
    try: threshold = float(threshold_row["value"] if threshold_row and threshold_row["value"] else 75)
    except Exception: threshold = 75
    rows = db.execute("""SELECT u.id, s.name, COUNT(a.id) total,
        COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present
        FROM users u JOIN students s ON s.id=u.student_id LEFT JOIN attendance a ON a.student_id=s.id
        WHERE u.role='student' GROUP BY u.id,s.name HAVING COUNT(a.id)>0""").fetchall()
    for r in rows:
        pct = 100.0*r["present"]/r["total"] if r["total"] else 0
        if pct < threshold:
            recent = db.execute("SELECT id FROM notifications WHERE user_id=? AND title='Low Attendance Alert' AND created_at>=? LIMIT 1",
                                (r["id"], datetime.utcnow().strftime('%Y-%m-01T00:00:00'))).fetchone()
            if not recent:
                notify_user(db, r["id"], "Low Attendance Alert", f"Your attendance is {pct:.1f}%, below the {threshold:.0f}% threshold.", "/student/dashboard")


@app.route("/notifications")
@login_required
def notifications():
    db=get_db()
    rows=db.execute("SELECT * FROM notifications WHERE user_id=? ORDER BY created_at DESC LIMIT 100", (session["user_id"],)).fetchall()
    db.execute("UPDATE notifications SET is_read=1 WHERE user_id=?", (session["user_id"],)); db.commit(); db.close()
    return render_template("notifications.html", notifications=rows)


@app.route("/notifications/read/<int:notification_id>")
@login_required
def read_notification(notification_id):
    db=get_db(); row=db.execute("SELECT link FROM notifications WHERE id=? AND user_id=?", (notification_id,session["user_id"])).fetchone()
    db.execute("UPDATE notifications SET is_read=1 WHERE id=? AND user_id=?", (notification_id,session["user_id"])); db.commit(); db.close()
    return redirect(row["link"] if row and row["link"] else url_for("notifications"))


@app.route("/departments", methods=["GET","POST"])
@login_required
@admin_required
def departments():
    db=get_db()
    if request.method=="POST":
        code=request.form.get("code","").strip().upper(); name=request.form.get("name","").strip()
        if not code or not name: flash("Department code and name are required.","error")
        else:
            try:
                db.execute("INSERT INTO departments(code,name,active) VALUES(?,?,1)",(code,name)); db.commit(); flash("Department added.","success")
            except Exception:
                db.conn.rollback(); flash("Department code or name already exists.","error")
    depts=db.execute("SELECT * FROM departments ORDER BY name").fetchall()
    teachers=db.execute("SELECT id,username,full_name,department_id FROM users WHERE role='teacher' AND approved=1 ORDER BY full_name").fetchall()
    db.close(); return render_template("departments.html", departments=depts, teachers=teachers)


@app.post("/departments/edit/<int:department_id>")
@login_required
@admin_required
def edit_department(department_id):
    code=request.form.get("code","").strip().upper(); name=request.form.get("name","").strip(); active=1 if request.form.get("active","1")=="1" else 0
    db=get_db()
    if not code or not name:
        flash("Department code and name are required.","error")
    else:
        try:
            db.execute("UPDATE departments SET code=?,name=?,active=? WHERE id=?",(code,name,active,department_id)); db.commit(); flash("Department updated.","success")
        except Exception:
            db.conn.rollback(); flash("Department code or name already exists.","error")
    db.close(); return redirect(url_for("departments"))


@app.post("/departments/delete/<int:department_id>")
@login_required
@admin_required
def delete_department(department_id):
    db=get_db()
    student_count=db.execute("SELECT COUNT(*) c FROM students WHERE department_id=?",(department_id,)).fetchone()["c"]
    subject_count=db.execute("SELECT COUNT(*) c FROM subjects WHERE department_id=?",(department_id,)).fetchone()["c"]
    user_count=db.execute("SELECT COUNT(*) c FROM users WHERE department_id=?",(department_id,)).fetchone()["c"]
    if student_count or subject_count or user_count:
        flash("Cannot delete a department that still has students, subjects or staff. Reassign them first.","error")
    else:
        db.execute("DELETE FROM departments WHERE id=?",(department_id,)); db.commit(); flash("Department removed.","success")
    db.close(); return redirect(url_for("departments"))


@app.post("/departments/assign-hod")
@login_required
@admin_required
def assign_hod():
    db=get_db(); teacher_id=request.form.get("teacher_id",type=int); dept_id=request.form.get("department_id",type=int)
    if teacher_id and dept_id:
        db.execute("UPDATE users SET role='teacher' WHERE role='hod' AND id<>? AND department_id=?", (teacher_id,dept_id))
        db.execute("UPDATE users SET role='hod',department_id=? WHERE id=? AND approved=1", (dept_id,teacher_id)); db.commit(); flash("HOD assigned successfully.","success")
    db.close(); return redirect(url_for("departments"))


@app.route("/hod/dashboard")
@login_required
def hod_dashboard():
    if session.get("role")!="hod": return redirect(url_for("dashboard"))
    db=get_db(); dept=db.execute("SELECT d.* FROM departments d JOIN users u ON u.department_id=d.id WHERE u.id=?",(session["user_id"],)).fetchone()
    dept_id=dept["id"] if dept else None
    students=db.execute("SELECT COUNT(*) c FROM students WHERE department_id=?",(dept_id,)).fetchone()["c"] if dept_id else 0
    teachers=db.execute("SELECT COUNT(*) c FROM users WHERE department_id=? AND role='teacher'",(dept_id,)).fetchone()["c"] if dept_id else 0
    low=[]
    threshold=75
    tr=db.execute("SELECT value FROM app_settings WHERE key='attendance_threshold'").fetchone()
    try: threshold=float(tr["value"]) if tr else 75
    except: threshold=75
    low=db.execute("""SELECT s.prn,s.name,s.year,COUNT(a.id) total,COALESCE(SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END),0) present,
        CASE WHEN COUNT(a.id)=0 THEN 0 ELSE ROUND(100.0*SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END)/COUNT(a.id),1) END pct
        FROM students s LEFT JOIN attendance a ON a.student_id=s.id WHERE s.department_id=? GROUP BY s.id ORDER BY pct""",(dept_id,)).fetchall() if dept_id else []
    low=[r for r in low if r["total"] and float(r["pct"])<threshold][:20]
    db.close(); return render_template("hod_dashboard.html",department=dept,students=students,teachers=teachers,low=low,threshold=threshold)


@app.route("/leaves", methods=["GET","POST"])
@login_required
def leaves():
    db=get_db()
    if session.get("role")=="student":
        st=current_student(db)
        if request.method=="POST":
            fd=request.form.get("from_date",""); td=request.form.get("to_date",""); reason=request.form.get("reason","").strip()
            if not fd or not td or not reason or fd>td: flash("Enter valid leave dates and reason.","error")
            else:
                db.execute("INSERT INTO leave_applications(student_id,from_date,to_date,reason,status,created_at) VALUES(?,?,?,?,?,?)",(st["id"],fd,td,reason,"Pending",datetime.utcnow().isoformat()))
                staff=db.execute("SELECT id FROM users WHERE role IN ('admin','hod','teacher') AND approved=1").fetchall()
                for x in staff: notify_user(db,x["id"],"New Leave Application",f"{st['name']} submitted a leave application.","/leaves")
                db.commit(); flash("Leave application submitted.","success")
        rows=db.execute("SELECT * FROM leave_applications WHERE student_id=? ORDER BY created_at DESC",(st["id"],)).fetchall(); db.close()
        return render_template("leaves.html", student_mode=True, leaves=rows)
    # staff view; HOD only sees own department
    if session.get("role")=='hod':
        rows=db.execute("""SELECT l.*,s.name,s.prn,s.year FROM leave_applications l JOIN students s ON s.id=l.student_id
            JOIN users u ON u.id=? WHERE s.department_id=u.department_id ORDER BY l.created_at DESC""",(session["user_id"],)).fetchall()
    else:
        rows=db.execute("SELECT l.*,s.name,s.prn,s.year FROM leave_applications l JOIN students s ON s.id=l.student_id ORDER BY l.created_at DESC").fetchall()
    db.close(); return render_template("leaves.html", student_mode=False, leaves=rows)


@app.post("/leaves/<int:leave_id>/review")
@login_required
def review_leave(leave_id):
    if session.get("role") not in ('admin','hod'):
        flash("Only Admin or HOD can review leave applications.","error"); return redirect(url_for("leaves"))
    status=request.form.get("status"); note=request.form.get("review_note","").strip()
    if status not in ('Approved','Rejected'): return redirect(url_for("leaves"))
    db=get_db(); row=db.execute("SELECT l.*,u.id user_id,s.name FROM leave_applications l JOIN students s ON s.id=l.student_id JOIN users u ON u.student_id=s.id WHERE l.id=?",(leave_id,)).fetchone()
    if row:
        db.execute("UPDATE leave_applications SET status=?,reviewed_by=?,review_note=? WHERE id=?",(status,session["user_id"],note,leave_id))
        notify_user(db,row["user_id"],"Leave Application Updated",f"Your leave application has been {status.lower()}.","/leaves")
        db.commit(); flash(f"Leave {status.lower()}.","success")
    db.close(); return redirect(url_for("leaves"))


@app.route("/qr-attendance", methods=["GET","POST"])
@login_required
@staff_required
def qr_attendance():
    db=get_db(); subjects=teacher_or_admin_subjects(db); years=YEAR_OPTIONS[:]
    selected_year=request.args.get('year') or request.form.get('year') or '1st Year'
    session_type=(request.args.get('session_type') or request.form.get('session_type') or 'Lecture').strip()
    if session_type not in ('Lecture','Practical'): session_type='Lecture'
    batch=(request.args.get('batch') or request.form.get('batch') or '').strip()
    subject_id=request.args.get('subject_id',type=int) or request.form.get('subject_id',type=int)
    ad=request.args.get('attendance_date') or request.form.get('attendance_date') or date.today().isoformat()
    tid=request.args.get('timetable_id',type=int) or request.form.get('timetable_id',type=int)
    start_time=request.args.get('start_time') or request.form.get('start_time',''); end_time=request.args.get('end_time') or request.form.get('end_time','')
    teacher_status=request.args.get('teacher_status') or request.form.get('teacher_status','Scheduled teacher present'); conducted_by=request.args.get('conducted_by',type=int) or request.form.get('conducted_by',type=int)
    subjects=[s for s in subjects if s['year']==selected_year]
    if session_type=='Practical' and batch not in ('Batch A','Batch B','Batch C','Batch D'): batch=''
    day_name=date.fromisoformat(ad).strftime('%A') if ad else date.today().strftime('%A')
    candidates=timetable_candidates(db,selected_year,session_type,batch,subject_id,day_name) if subject_id else []; teachers=all_teacher_rows(db)
    if request.method=='POST':
        tr=db.execute('SELECT * FROM timetable WHERE id=?',(tid,)).fetchone() if tid else None
        if not subject_id or not any(s['id']==subject_id for s in subjects) or not start_time or not end_time:
            flash('Select valid timetable, subject and start/end time.','error')
        elif session_type=='Practical' and batch not in ('Batch A','Batch B','Batch C','Batch D'):
            flash('Select a valid practical batch.','error')
        else:
            scheduled=tr['teacher_id'] if tr else None
            if teacher_status=='Scheduled teacher present': conducted_by=scheduled or session['user_id']
            elif not conducted_by: conducted_by=session['user_id']
            token=secrets.token_urlsafe(24); exp=datetime.utcnow().replace(microsecond=0)+__import__('datetime').timedelta(minutes=10)
            db.execute("""INSERT INTO qr_sessions(token,subject_id,year,attendance_date,lecture_no,lecture_time,start_time,end_time,session_type,batch,teacher_id,scheduled_teacher_id,conducted_by_teacher_id,teacher_status,expires_at,active)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)""",(token,subject_id,selected_year,ad,1,f'{start_time}-{end_time}',start_time,end_time,session_type,batch if session_type=='Practical' else '',session['user_id'],scheduled,conducted_by,teacher_status,exp.isoformat()))
            db.commit(); db.close(); return redirect(url_for('qr_attendance',token=token))
    token=request.args.get('token',''); qr=None; qr_session=None
    if token:
        qr_session=db.execute('SELECT q.*,s.code,s.name subject_name FROM qr_sessions q JOIN subjects s ON s.id=q.subject_id WHERE q.token=? AND q.active=1',(token,)).fetchone()
        if qr_session and datetime.fromisoformat(qr_session['expires_at'])<datetime.utcnow(): qr_session=None
        if qr_session and QRCODE_AVAILABLE:
            import base64
            img=qrcode.make(request.url_root.rstrip('/')+url_for('qr_scan',token=token)); bio=io.BytesIO(); img.save(bio,format='PNG'); qr='data:image/png;base64,'+base64.b64encode(bio.getvalue()).decode()
    db.close(); return render_template('qr_attendance.html',subjects=subjects,years=years,today=date.today().isoformat(),qr=qr,qr_session=qr_session,token=token,batch=batch,session_type=session_type,candidates=candidates,teachers=teachers,selected_year=selected_year,day_name=day_name)


@app.route('/qr/scan',methods=['GET','POST'])
@login_required
def qr_scan():
    if session.get('role')!='student': flash('Only students can scan attendance QR codes.','error'); return redirect(url_for('dashboard'))
    token=request.args.get('token') or request.form.get('token',''); db=get_db()
    row=db.execute("""SELECT q.*,s.code,s.name subject_name FROM qr_sessions q JOIN subjects s ON s.id=q.subject_id WHERE q.token=? AND q.active=1""",(token,)).fetchone()
    if not row or datetime.fromisoformat(row['expires_at'])<datetime.utcnow(): db.close(); return render_template('qr_scan.html',valid=False)
    st=current_student(db)
    if st['year']!=row['year'] or (row['session_type']=='Practical' and st['batch']!=row['batch']) or (st['department_id'] and db.execute('SELECT department_id FROM subjects WHERE id=?',(row['subject_id'],)).fetchone()['department_id']!=st['department_id']): db.close(); return render_template('qr_scan.html',valid=False,reason='This QR is not for your class or practical batch.')
    if request.method=='POST':
        db.execute("""INSERT INTO attendance(student_id,subject_id,attendance_date,lecture_no,lecture_time,session_type,batch,status,marked_by) VALUES(?,?,?,?,?,?,?,?,?)
            ON CONFLICT(student_id,subject_id,attendance_date,lecture_no) DO UPDATE SET lecture_time=excluded.lecture_time,status=excluded.status,marked_by=excluded.marked_by""",
                   (st['id'],row['subject_id'],row['attendance_date'],row['lecture_no'],row['lecture_time'],row['session_type'],row['batch'],'Present',row['teacher_id']))
        notify_user(db, session['user_id'], "Attendance Marked", f"You were marked Present for {row['subject_name']} (Lecture {row['lecture_no']}).", "/student/dashboard")
        create_low_attendance_notifications(db)
        db.commit(); db.close(); return render_template('qr_scan.html',valid=True,done=True,session_info=row)
    db.close(); return render_template('qr_scan.html',valid=True,session_info=row)


@app.post('/qr/close/<token>')
@login_required
@staff_required
def close_qr(token):
    db=get_db(); db.execute("UPDATE qr_sessions SET active=0 WHERE token=? AND teacher_id=?",(token,session['user_id'])); db.commit(); db.close(); flash('QR attendance session closed.','success'); return redirect(url_for('qr_attendance'))


@app.get('/manifest.json')
def manifest():
    return {"name":COLLEGE_NAME+" Attendance","short_name":"YNP Attendance","start_url":"/","display":"standalone","background_color":"#ffffff","theme_color":"#1f4e79","icons":[{"src":url_for('static',filename=LOGO_FILE),"sizes":"192x192","type":"image/png"},{"src":url_for('static',filename=LOGO_FILE),"sizes":"512x512","type":"image/png"}]}


@app.get('/service-worker.js')
def service_worker():
    js="""self.addEventListener('install',e=>self.skipWaiting());self.addEventListener('activate',e=>self.clients.claim());self.addEventListener('notificationclick',e=>{e.notification.close();e.waitUntil(clients.openWindow('/notifications'))});"""
    return Response(js,mimetype='application/javascript')

# Initialize the database when the application is imported by Gunicorn/Render.
init_db()

if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG", "0") == "1",
            host=os.environ.get("HOST", "127.0.0.1"),
            port=int(os.environ.get("PORT", "5000")))
